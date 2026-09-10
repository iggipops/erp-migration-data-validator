from openpyxl.styles import PatternFill
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from config.config_loader import AppConfig, ISSUES_SHEET
from workbook.excel_utils import remove_sheet_if_exists
from utils.logger import get_logger


ISSUES_HEADERS = [
    "IssueNum",
    "SheetName",
    "ColumnName",
    "RowNumber",
    "CellContent",
    "RuleCode",
    "RuleName",
    "Severity",
    "Comment",
]


class IssuesWriter:

    def __init__(self, workbook: Workbook, config: AppConfig):
        self.workbook      = workbook
        self.config        = config
        self.logger        = get_logger()

        self.issue_counter   = 0
        self.error_counter   = 0
        self.warning_counter = 0

        self._sheet: Worksheet = self._initialize_sheet()

    # -----------------------------------------------------------------------

    def record(
        self,
        cell,
        rule: dict,
        sheet_name: str,
        column_name: str,
        row_number: int,
        cell_content,
        comment: str = "",
    ) -> None:
        """Highlight the cell and append a row to the Issues sheet."""

        # Highlight cell
        cell.fill = self._get_fill(str(rule.get("Color", "")).strip())

        self.issue_counter += 1
        severity = str(rule.get("Severity", "")).strip().upper()

        if severity == "ERROR":
            self.error_counter += 1
        elif severity == "WARNING":
            self.warning_counter += 1

        self._sheet.append([
            self.issue_counter,
            sheet_name,
            column_name,
            row_number,
            cell_content,
            rule.get("RuleCode", ""),
            rule.get("RuleName", ""),
            rule.get("Severity", ""),
            comment or "",
        ])

    # -----------------------------------------------------------------------

    def _initialize_sheet(self) -> Worksheet:
        remove_sheet_if_exists(self.workbook, ISSUES_SHEET)
        sheet = self.workbook.create_sheet(ISSUES_SHEET)
        sheet.append(ISSUES_HEADERS)

        # Apply autofilter to header row
        sheet.auto_filter.ref = f"A1:{chr(64 + len(ISSUES_HEADERS))}1"

        self.logger.debug("Issues worksheet initialized")
        return sheet

    def apply_column_widths(self) -> None:
        """Call after all issues are written to auto-fit column widths."""
        col_widths = {i + 1: len(h) for i, h in enumerate(ISSUES_HEADERS)}

        for row in self._sheet.iter_rows(min_row=2, values_only=True):
            for i, value in enumerate(row):
                if value is not None:
                    col_widths[i + 1] = max(
                        col_widths.get(i + 1, 0),
                        min(len(str(value)), 60)  # cap at 60 chars
                    )

        for col_idx, width in col_widths.items():
            col_letter = chr(64 + col_idx)
            self._sheet.column_dimensions[col_letter].width = width + 2

    def _get_fill(self, color_value: str) -> PatternFill:
        upper = color_value.upper()
        rgb   = self.config.color_map.get(upper, upper)
        return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")
