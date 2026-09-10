from datetime import datetime

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook import Workbook

from config.config_loader import AppConfig, SUMMARY_SHEET, ISSUES_SHEET
from output.issues_writer import IssuesWriter
from workbook.excel_utils import remove_sheet_if_exists, sheet_exists, get_headers
from utils.logger import get_logger


# Styling
HEADER_FILL  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
TOTAL_FILL   = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
SECTION_FILL = PatternFill(start_color="E9EFF7", end_color="E9EFF7", fill_type="solid")
HEADER_FONT  = Font(bold=True, color="FFFFFF")
TOTAL_FONT   = Font(bold=True)
SECTION_FONT = Font(bold=True, size=11)


class SummaryWriter:

    def __init__(
        self,
        workbook: Workbook,
        config: AppConfig,
        issues_writer: IssuesWriter,
    ):
        self.workbook      = workbook
        self.config        = config
        self.issues_writer = issues_writer
        self.logger        = get_logger()

    def write(
        self,
        enrichment_stats: list[dict],
        validation_stats: list[dict],
        external_files_result: dict,
        external_files_rules: list[dict],
        start_time: datetime,
    ) -> None:

        remove_sheet_if_exists(self.workbook, SUMMARY_SHEET)
        sheet = self.workbook.create_sheet(SUMMARY_SHEET, 0)

        row = 1

        # --- Run Information ---
        row = self._write_section_title(sheet, row, "Run Information")
        row = self._write_kv(sheet, row, "Timestamp",   start_time.strftime("%Y-%m-%d %H:%M:%S"))
        row = self._write_kv(sheet, row, "Input file",  self.config.input_file)
        row = self._write_kv(sheet, row, "Output file", self.config.output_file)
        row += 1

        # --- External Files Statistics ---
        row = self._write_section_title(sheet, row, "External Files Statistics")
        row = self._write_external_files_table(
            sheet, row, external_files_rules, external_files_result
        )
        row += 1

        # --- Enrichment Rules Statistics ---
        row = self._write_section_title(sheet, row, "Enrichment Rules Statistics")
        row = self._write_enrichment_table(sheet, row, enrichment_stats)
        row += 1

        # --- Validation Rules Statistics ---
        row = self._write_section_title(sheet, row, "Validation Rules Statistics")
        row = self._write_validation_table(sheet, row, validation_stats)
        row += 1

        # --- Validation Issues Statistics by Rules ---
        row = self._write_section_title(sheet, row, "Validation Issues Statistics by Rules")
        row = self._write_issues_by_rules_table(sheet, row, validation_stats)
        row += 1

        # --- Validation Issues Statistics by Sheet/Column ---
        row = self._write_section_title(sheet, row, "Validation Issues Statistics by Sheet / Column")
        row = self._write_issues_by_sheet_table(sheet, row, validation_stats)

        # Column widths
        widths = [12, 35, 35, 20, 20, 12, 12, 12]
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + i)].width = w

        self.logger.info("Summary worksheet written")

    def write_metadata_errors(
        self,
        errors: list[str],
        start_time: datetime,
        external_files_result: dict | None = None,
        external_files_rules: list[dict] | None = None,
    ) -> None:
        """
        FS section 8.6: when ValidationRules/EnrichmentRules metadata
        validation fails, execution terminates before any enrichment or
        validation runs — but the Summary worksheet must still be written,
        listing every error, so the output file explains the termination.

        Step 5 (external file import) already ran before Step 6 terminates,
        so when its result is available it's included here the same way
        write() includes it, rather than being silently dropped.
        """
        remove_sheet_if_exists(self.workbook, SUMMARY_SHEET)
        sheet = self.workbook.create_sheet(SUMMARY_SHEET, 0)

        row = 1

        row = self._write_section_title(sheet, row, "Run Information")
        row = self._write_kv(sheet, row, "Timestamp",   start_time.strftime("%Y-%m-%d %H:%M:%S"))
        row = self._write_kv(sheet, row, "Input file",  self.config.input_file)
        row = self._write_kv(sheet, row, "Output file", self.config.output_file)
        row = self._write_kv(sheet, row, "Result",      "TERMINATED — metadata validation failed")
        row += 1

        if external_files_result is not None:
            row = self._write_section_title(sheet, row, "External Files Statistics")
            row = self._write_external_files_table(
                sheet, row, external_files_rules or [], external_files_result
            )
            row += 1

        row = self._write_section_title(sheet, row, "Metadata Validation Errors")
        row = self._write_table_header(sheet, row, ["#", "Error"])
        for i, err in enumerate(errors, start=1):
            sheet.cell(row=row, column=1, value=i)
            sheet.cell(row=row, column=2, value=err)
            row += 1

        total_label = f"Metadata Validation Errors: {len(errors)}"
        row = self._write_total_row(sheet, row, total_label, col_span=2)

        widths = [12, 35, 35, 20, 20, 12, 12, 12]
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + i)].width = w

        self.logger.info("Summary worksheet written (metadata validation failure)")

    # -----------------------------------------------------------------------
    # External Files table
    # -----------------------------------------------------------------------

    def _write_external_files_table(
        self, sheet, row: int,
        rules: list[dict], result: dict
    ) -> int:

        headers = ["SequenceNum", "FilePath", "Status", "ConversionWarnings"]
        row = self._write_table_header(sheet, row, headers)

        active_rules = [
            r for r in rules
            if str(r.get("Active", "")).strip().upper() == "YES"
        ]

        loaded = set(result.get("loaded", []))
        failed = set(result.get("failed_sheets") or set())
        if not failed:
            # Back-compat: derive sheet names from failure messages when
            # failed_sheets wasn't provided (message format: "'<sheet>' from '<path>': <reason>").
            failed = {msg.split("'")[1] for msg in result.get("failed", []) if "'" in msg}
        conv_warns = result.get("conversion_warnings", {})

        total_active   = len(active_rules)
        total_imported = 0
        total_skipped  = 0
        total_failed   = 0
        total_warnings = 0

        for r in active_rules:
            seq       = r.get("SequenceNum", "")
            file_path = str(r.get("FilePath", "")).strip()
            sheet_name = str(r.get("SheetName", "")).strip()

            if sheet_name in loaded:
                status = "imported"
                total_imported += 1
            elif file_path in failed or sheet_name in failed:
                status = "failed"
                total_failed += 1
            else:
                status = "skipped"
                total_skipped += 1

            warn_count = conv_warns.get(sheet_name, 0)
            total_warnings += warn_count

            for col_idx, value in enumerate([seq, file_path, status, warn_count], start=1):
                sheet.cell(row=row, column=col_idx, value=value)
            row += 1

        # Total row
        total_label = (
            f"Active External Files: {total_active}     "
            f"imported: {total_imported}     "
            f"skipped: {total_skipped}     "
            f"failed: {total_failed}     "
            f"conversion warnings: {total_warnings}"
        )
        row = self._write_total_row(sheet, row, total_label, col_span=4)
        return row

    # -----------------------------------------------------------------------
    # Enrichment Rules table
    # -----------------------------------------------------------------------

    def _write_enrichment_table(self, sheet, row: int, stats: list[dict]) -> int:

        headers = ["SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName", "Status"]
        row = self._write_table_header(sheet, row, headers)

        totals = {"executed": 0, "skipped": 0, "failed": 0}

        for stat in stats:
            status = stat.get("status", "")
            values = [
                stat.get("SequenceNum", ""),
                stat.get("RuleCode", ""),
                stat.get("RuleName", ""),
                stat.get("SheetName", ""),
                stat.get("ColumnName", ""),
                status,
            ]
            for col_idx, value in enumerate(values, start=1):
                sheet.cell(row=row, column=col_idx, value=value)
            row += 1
            totals[status] = totals.get(status, 0) + 1

        total_active = sum(totals.values())
        total_label = (
            f"Active Enrichment Rules: {total_active}     "
            f"executed: {totals['executed']}     "
            f"skipped: {totals['skipped']}     "
            f"failed: {totals['failed']}"
        )
        row = self._write_total_row(sheet, row, total_label, col_span=6)
        return row

    # -----------------------------------------------------------------------
    # Validation Rules table
    # -----------------------------------------------------------------------

    def _write_validation_table(self, sheet, row: int, stats: list[dict]) -> int:

        headers = ["SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName",
                   "Status", "RowsUnchecked"]
        row = self._write_table_header(sheet, row, headers)

        totals = {"executed": 0, "skipped": 0, "failed": 0, "partial": 0}

        for stat in stats:
            status = stat.get("status", "")
            values = [
                stat.get("SequenceNum", ""),
                stat.get("RuleCode", ""),
                stat.get("RuleName", ""),
                stat.get("SheetName", ""),
                stat.get("ColumnName", ""),
                status,
                stat.get("rows_unchecked", 0),
            ]
            for col_idx, value in enumerate(values, start=1):
                sheet.cell(row=row, column=col_idx, value=value)
            row += 1
            totals[status] = totals.get(status, 0) + 1

        total_active = sum(totals.values())
        total_label = (
            f"Active Validation Rules: {total_active}     "
            f"executed: {totals['executed']}     "
            f"skipped: {totals['skipped']}     "
            f"failed: {totals['failed']}     "
            f"partial: {totals['partial']}"
        )
        row = self._write_total_row(sheet, row, total_label, col_span=7)
        return row

    # -----------------------------------------------------------------------
    # Issues by Rules table
    # -----------------------------------------------------------------------

    def _write_issues_by_rules_table(self, sheet, row: int, stats: list[dict]) -> int:

        headers = ["SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName", "Severity", "Number of Issues"]
        row = self._write_table_header(sheet, row, headers)

        issue_counts     = self._count_issues_per_rule()
        # partial AI rules (FS section 6.8) DO have real issues recorded for
        # their successful-batch rows — include them alongside executed rules.
        executed_stats   = [s for s in stats if s.get("status") in ("executed", "partial")]

        total_issues   = 0
        total_errors   = 0
        total_warnings = 0

        for stat in executed_stats:
            rule_code = stat.get("RuleCode", "")
            counts    = issue_counts.get(rule_code, {"ERROR": 0, "WARNING": 0})
            errors    = counts.get("ERROR", 0)
            warnings  = counts.get("WARNING", 0)
            issues    = errors + warnings
            severity  = stat.get("Severity", "")

            values = [
                stat.get("SequenceNum", ""),
                rule_code,
                stat.get("RuleName", ""),
                stat.get("SheetName", ""),
                stat.get("ColumnName", ""),
                severity,
                issues,
            ]
            for col_idx, value in enumerate(values, start=1):
                sheet.cell(row=row, column=col_idx, value=value)
            row += 1

            total_issues   += issues
            total_errors   += errors
            total_warnings += warnings

        total_label = (
            f"Active Executed/Partial Validation Rules: {len(executed_stats)}     "
            f"Issues found: {total_issues}     "
            f"Errors found: {total_errors}     "
            f"Warnings found: {total_warnings}"
        )
        row = self._write_total_row(sheet, row, total_label, col_span=7)
        return row

    # -----------------------------------------------------------------------
    # Issues by Sheet/Column table
    # -----------------------------------------------------------------------

    def _write_issues_by_sheet_table(self, sheet, row: int, stats: list[dict]) -> int:

        headers = ["SheetName", "ColumnName", "Number of Errors", "Number of Warnings"]
        row = self._write_table_header(sheet, row, headers)

        # Build sheet/column -> counts from Issues sheet
        sheet_col_counts = self._count_issues_per_sheet_column()
        executed_stats   = [s for s in stats if s.get("status") in ("executed", "partial")]

        # Collect unique sheet/column pairs in execution order
        seen   = set()
        pairs  = []
        for stat in executed_stats:
            key = (stat.get("SheetName", ""), stat.get("ColumnName", ""))
            if key not in seen:
                seen.add(key)
                pairs.append(key)

        total_errors   = 0
        total_warnings = 0

        for sheet_name, col_name in pairs:
            key      = (sheet_name, col_name)
            counts   = sheet_col_counts.get(key, {"ERROR": 0, "WARNING": 0})
            errors   = counts.get("ERROR", 0)
            warnings = counts.get("WARNING", 0)

            for col_idx, value in enumerate(
                [sheet_name, col_name, errors, warnings], start=1
            ):
                sheet.cell(row=row, column=col_idx, value=value)
            row += 1

            total_errors   += errors
            total_warnings += warnings

        total_issues = total_errors + total_warnings
        total_label = (
            f"Active Executed/Partial Validation Rules: {len(executed_stats)}     "
            f"Issues found: {total_issues}     "
            f"Errors found: {total_errors}     "
            f"Warnings found: {total_warnings}"
        )
        row = self._write_total_row(sheet, row, total_label, col_span=4)
        return row

    # -----------------------------------------------------------------------
    # Helpers
    # -----------------------------------------------------------------------

    def _write_section_title(self, sheet, row: int, title: str) -> int:
        cell      = sheet.cell(row=row, column=1, value=title)
        cell.font = SECTION_FONT
        cell.fill = SECTION_FILL
        return row + 1

    def _write_table_header(self, sheet, row: int, headers: list[str]) -> int:
        for col_idx, col_name in enumerate(headers, start=1):
            cell      = sheet.cell(row=row, column=col_idx, value=col_name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        return row + 1

    def _write_total_row(self, sheet, row: int, label: str, col_span: int) -> int:
        cell = sheet.cell(row=row, column=1, value=label)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        # Merge across all columns
        if col_span > 1:
            sheet.merge_cells(
                start_row=row, start_column=1,
                end_row=row, end_column=col_span
            )
        return row + 1

    def _write_kv(self, sheet, row: int, key: str, value) -> int:
        sheet.cell(row=row, column=1, value=key)
        sheet.cell(row=row, column=2, value=str(value))
        return row + 1

    def _count_issues_per_rule(self) -> dict[str, dict[str, int]]:
        counts: dict[str, dict[str, int]] = {}
        if not sheet_exists(self.workbook, ISSUES_SHEET):
            return counts

        ws      = self.workbook[ISSUES_SHEET]
        headers = get_headers(ws)
        rc_idx  = headers.get("RuleCode")
        sev_idx = headers.get("Severity")

        if not rc_idx or not sev_idx:
            return counts

        for row in range(2, ws.max_row + 1):
            rc  = str(ws.cell(row=row, column=rc_idx).value or "")
            sev = str(ws.cell(row=row, column=sev_idx).value or "").upper()
            if not rc:
                continue
            if rc not in counts:
                counts[rc] = {"ERROR": 0, "WARNING": 0}
            if sev in ("ERROR", "WARNING"):
                counts[rc][sev] += 1

        return counts

    def _count_issues_per_sheet_column(self) -> dict[tuple, dict[str, int]]:
        counts: dict[tuple, dict[str, int]] = {}
        if not sheet_exists(self.workbook, ISSUES_SHEET):
            return counts

        ws      = self.workbook[ISSUES_SHEET]
        headers = get_headers(ws)
        sn_idx  = headers.get("SheetName")
        cn_idx  = headers.get("ColumnName")
        sev_idx = headers.get("Severity")

        if not sn_idx or not cn_idx or not sev_idx:
            return counts

        for row in range(2, ws.max_row + 1):
            sn  = str(ws.cell(row=row, column=sn_idx).value or "")
            cn  = str(ws.cell(row=row, column=cn_idx).value or "")
            sev = str(ws.cell(row=row, column=sev_idx).value or "").upper()
            key = (sn, cn)
            if not sn:
                continue
            if key not in counts:
                counts[key] = {"ERROR": 0, "WARNING": 0}
            if sev in ("ERROR", "WARNING"):
                counts[key][sev] += 1

        return counts
