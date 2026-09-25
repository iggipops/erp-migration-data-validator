import re
from openpyxl.workbook import Workbook

from config.config_loader import (
    AppConfig,
    SUPPORTED_VALIDATION_TYPES,
    SUPPORTED_ENRICHMENT_TYPES,
    SUPPORTED_SEVERITIES,
    VALIDATION_RULES_SHEET,
    ENRICHMENT_RULES_SHEET,
)
from workbook.excel_utils import (
    sheet_exists, is_empty, normalize_string, clean_field,
    get_headers, get_sheet, find_column_index, cell_value, last_data_row,
)
from utils.logger import get_logger


# ---------------------------------------------------------------------------
# Required columns per metadata sheet
# ---------------------------------------------------------------------------

MANDATORY_VALIDATION_COLUMNS = [
    "SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName",
    "ValidationType", "Severity", "Color", "Active",
    "ReferenceDataSheet", "ReferenceDataColumn", "CustomFunctionName",
]

MANDATORY_ENRICHMENT_COLUMNS = [
    "SequenceNum", "RuleCode", "RuleName", "SheetName", "TargetColumn",
    "CustomFunctionName", "FunctionArguments", "Active",
]

CONDITIONAL_ENRICHMENT_COLUMNS = [
    "Color",
]

class MetadataValidator:

    def __init__(
        self, workbook: Workbook, config: AppConfig, registry, vtype_registry=None,
        failed_sheets: set[str] | None = None,
    ):
        self.workbook       = workbook
        self.config         = config
        self.registry       = registry
        self.vtype_registry = vtype_registry
        self.errors         = []
        self.config   = config
        self.registry = registry   # FunctionRegistry — for CustomFunctionName checks
        self.logger   = get_logger()
        self.errors: list[str] = []
        # SheetName values that failed external-file import (FS 8.5.4) —
        # rows depending on these are skipped, not raised as metadata errors.
        self.failed_sheets = {normalize_string(s) for s in (failed_sheets or set())}

    def validate(self) -> list[str]:
        """
        Validate all metadata sheets.
        Returns list of error strings (empty = all OK).
        """
        self.errors = []

        self._validate_validation_rules()
        self._validate_enrichment_rules()

        if self.errors:
            for err in self.errors:
                self.logger.error(f"Metadata error: {err}")
        else:
            self.logger.info("Metadata validation passed")

        return self.errors

    # -----------------------------------------------------------------------
    # ValidationRules
    # -----------------------------------------------------------------------

    def _validate_validation_rules(self) -> None:

        if not sheet_exists(self.workbook, VALIDATION_RULES_SHEET):
            self.errors.append(f'Sheet "{VALIDATION_RULES_SHEET}" not found')
            return

        sheet, headers = self._sheet_and_headers(VALIDATION_RULES_SHEET)

        header_error_count = 0
        for col in MANDATORY_VALIDATION_COLUMNS:
            if find_column_index(headers, col) is None:
                self.errors.append(
                    f'{VALIDATION_RULES_SHEET}: mandatory column "{col}" not found'
                )
                header_error_count += 1

        if header_error_count:
            return

        seq_nums, rule_codes, rule_names = [], [], []

        for row in range(2, last_data_row(sheet) + 1):

            r = self._row_dict(sheet, headers, row, MANDATORY_VALIDATION_COLUMNS)

            if self._all_empty(r):
                continue

            active = str(r.get("Active", "")).strip().upper()
            if active not in ("YES", "NO"):
                self.errors.append(
                    f'{VALIDATION_RULES_SHEET} row {row}: '
                    f'Active must be YES or NO, got "{active}"'
                )
                continue

            if active == "NO":
                continue

            # --- Skip rows depending on a SheetName that failed to import (FS 8.5.4) ---
            row_sheet_name = str(r.get("SheetName", "")).strip()
            if normalize_string(row_sheet_name) in self.failed_sheets:
                self.logger.warning(
                    f'{VALIDATION_RULES_SHEET} row {row}: row skipped: depends on '
                    f'SheetName "{row_sheet_name}" which failed to import'
                )
                continue

            # --- SequenceNum ---
            if is_empty(r["SequenceNum"]):
                self.errors.append(f'{VALIDATION_RULES_SHEET} row {row}: SequenceNum is empty')
            else:
                try:
                    seq_nums.append(int(r["SequenceNum"]))
                except (ValueError, TypeError):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'SequenceNum must be numeric'
                    )

            # --- RuleCode / RuleName ---
            self._check_required(r, row, VALIDATION_RULES_SHEET, "RuleCode", rule_codes)
            self._check_required(r, row, VALIDATION_RULES_SHEET, "RuleName", rule_names)

            # --- SheetName / ColumnName exist in workbook ---
            sheet_name = str(r.get("SheetName", "")).strip()
            col_name   = str(r.get("ColumnName", "")).strip()

            if sheet_name and not sheet_exists(self.workbook, sheet_name):
                self.errors.append(
                    f'{VALIDATION_RULES_SHEET} row {row}: '
                    f'SheetName "{sheet_name}" not found in workbook'
                )
            elif sheet_name and col_name:
                ws = get_sheet(self.workbook, sheet_name)
                if find_column_index(get_headers(ws), col_name) is None:
                    # Skip error if this column is a TargetColumn of an active EnrichmentRule
                    # — it will be created during enrichment execution before validation runs
                    if not self._is_enrichment_target(sheet_name, col_name):
                        self.errors.append(
                            f'{VALIDATION_RULES_SHEET} row {row}: '
                            f'ColumnName "{col_name}" not found in sheet "{sheet_name}"'
                        )

            # --- ValidationType ---
            vtype = str(r.get("ValidationType", "")).strip().upper()
            # Use the registry as authoritative source if available,
            # fall back to SUPPORTED_VALIDATION_TYPES constant
            if self.vtype_registry:
                type_valid = self.vtype_registry.has(vtype)
            else:
                type_valid = vtype in SUPPORTED_VALIDATION_TYPES
            if not type_valid:
                self.errors.append(
                    f'{VALIDATION_RULES_SHEET} row {row}: '
                    f'ValidationType "{vtype}" is not supported'
                )

            # --- Type-specific metadata check via validation type registry ---
            if self.vtype_registry:
                check_fn = self.vtype_registry.get_metadata_check_fn(vtype)
                if check_fn:
                    extra_errors = check_fn(r, self.workbook, self.config)
                    for err in (extra_errors or []):
                        self.errors.append(
                            f'{VALIDATION_RULES_SHEET} row {row}: {err}'
                        )

            # --- Severity ---
            severity = str(r.get("Severity", "")).strip().upper()
            if severity not in SUPPORTED_SEVERITIES:
                self.errors.append(
                    f'{VALIDATION_RULES_SHEET} row {row}: '
                    f'Severity "{severity}" must be ERROR or WARNING'
                )

            # --- Color (required for ValidationRules, FS 8.7.3) ---
            self._check_color(r, row, VALIDATION_RULES_SHEET, required=True)

            # --- CustomFunctionName must exist in registry ---
            if vtype == "CUSTOM":
                fn_name = str(r.get("CustomFunctionName", "")).strip()
                if is_empty(fn_name):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'CustomFunctionName is required for CUSTOM type'
                    )
                elif not self.registry.has(fn_name):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'CustomFunctionName "{fn_name}" not found in function registry'
                    )
                else:
                    # Check REQUIRED_PARAMS against config.yaml custom_functions block
                    required = self.registry.get_required_params(fn_name)
                    if required:
                        provided = set(
                            self.config.custom_functions.get(fn_name.strip().lower(), {}).keys()
                        )
                        missing = [p for p in required if p not in provided]
                        if missing:
                            self.errors.append(
                                f'{VALIDATION_RULES_SHEET} row {row}: '
                                f'config.yaml custom_functions.{fn_name.strip().lower()} '
                                f'is missing required parameter(s): {", ".join(missing)}'
                            )

                    # Function-specific metadata check hook (FS section 2.6.2)
                    check_fn = self.registry.get_metadata_check_fn(fn_name)
                    if check_fn:
                        extra_errors = check_fn(r, self.workbook, self.config)
                        for err in (extra_errors or []):
                            self.errors.append(
                                f'{VALIDATION_RULES_SHEET} row {row}: {err}'
                            )

            # --- AI: ReferenceDataSheet / ReferenceDataColumn (optional, no lookup role,
            # FS section 6.8/8.6.3) and CustomFunctionName (required only when
            # ai_enabled=true — a soft-skipped rule when ai_enabled=false is not
            # held to this check) ---
            if vtype == "AI":
                ref_sheet  = clean_field(r.get("ReferenceDataSheet"))
                ref_column = clean_field(r.get("ReferenceDataColumn"))

                # --- Skip rows depending on a ReferenceDataSheet that failed to import ---
                if not is_empty(ref_sheet) and normalize_string(ref_sheet) in self.failed_sheets:
                    self.logger.warning(
                        f'{VALIDATION_RULES_SHEET} row {row}: row skipped: depends on '
                        f'ReferenceDataSheet "{ref_sheet}" which failed to import'
                    )
                    continue

                # --- Consistency: ReferenceDataColumn must be empty when
                # ReferenceDataSheet is empty (AI-specific; REFERENCE/DUPLICATE2
                # keep their own independent ReferenceDataColumn requiredness) ---
                if is_empty(ref_sheet):
                    if not is_empty(ref_column):
                        self.errors.append(
                            f'{VALIDATION_RULES_SHEET} row {row}: '
                            f'ReferenceDataColumn "{ref_column}" must be empty when '
                            f'ReferenceDataSheet is empty for AI type'
                        )
                elif not sheet_exists(self.workbook, ref_sheet):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'ReferenceDataSheet "{ref_sheet}" not found in workbook'
                    )
                elif normalize_string(ref_sheet) == normalize_string(sheet_name):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'ReferenceDataSheet "{ref_sheet}" must differ from '
                        f'SheetName for an AI rule'
                    )
                elif not is_empty(ref_column):
                    ws = get_sheet(self.workbook, ref_sheet)
                    if find_column_index(get_headers(ws), ref_column) is None:
                        if not self._is_enrichment_target(ref_sheet, ref_column):
                            self.errors.append(
                                f'{VALIDATION_RULES_SHEET} row {row}: '
                                f'ReferenceDataColumn "{ref_column}" not found '
                                f'in sheet "{ref_sheet}"'
                            )

                # --- CustomFunctionName required only when ai_enabled=true ---
                if self.config.ai_enabled:
                    fn_name = str(r.get("CustomFunctionName", "")).strip()
                    if is_empty(fn_name):
                        self.errors.append(
                            f'{VALIDATION_RULES_SHEET} row {row}: '
                            f'CustomFunctionName is required for AI type when ai_enabled=true'
                        )
                    elif not self.registry.has(fn_name):
                        self.errors.append(
                            f'{VALIDATION_RULES_SHEET} row {row}: '
                            f'CustomFunctionName "{fn_name}" not found in function registry'
                        )
                    else:
                        required = self.registry.get_required_params(fn_name)
                        if required:
                            provided = set(
                                self.config.custom_functions.get(fn_name.strip().lower(), {}).keys()
                            )
                            missing = [p for p in required if p not in provided]
                            if missing:
                                self.errors.append(
                                    f'{VALIDATION_RULES_SHEET} row {row}: '
                                    f'config.yaml custom_functions.{fn_name.strip().lower()} '
                                    f'is missing required parameter(s): {", ".join(missing)}'
                                )

                        check_fn = self.registry.get_metadata_check_fn(fn_name)
                        if check_fn:
                            extra_errors = check_fn(r, self.workbook, self.config)
                            for err in (extra_errors or []):
                                self.errors.append(
                                    f'{VALIDATION_RULES_SHEET} row {row}: {err}'
                                )

            # --- REFERENCE: ReferenceDataSheet / ReferenceDataColumn ---
            if vtype == "REFERENCE":
                ref_sheet  = clean_field(r.get("ReferenceDataSheet"))
                ref_column = clean_field(r.get("ReferenceDataColumn"))

                # --- Skip rows depending on a ReferenceDataSheet that failed to import (FS 8.5.4) ---
                if normalize_string(ref_sheet) in self.failed_sheets:
                    self.logger.warning(
                        f'{VALIDATION_RULES_SHEET} row {row}: row skipped: depends on '
                        f'ReferenceDataSheet "{ref_sheet}" which failed to import'
                    )
                    continue

                if is_empty(ref_sheet):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'ReferenceDataSheet is required for REFERENCE type'
                    )
                elif not sheet_exists(self.workbook, ref_sheet):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'ReferenceDataSheet "{ref_sheet}" not found in workbook'
                    )
                elif normalize_string(ref_sheet) == normalize_string(sheet_name):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'ReferenceDataSheet "{ref_sheet}" must differ from '
                        f'SheetName for a REFERENCE rule'
                    )
                elif not is_empty(ref_column):
                    ws = get_sheet(self.workbook, ref_sheet)
                    if find_column_index(get_headers(ws), ref_column) is None:
                        if not self._is_enrichment_target(ref_sheet, ref_column):
                            self.errors.append(
                                f'{VALIDATION_RULES_SHEET} row {row}: '
                                f'ReferenceDataColumn "{ref_column}" not found '
                                f'in sheet "{ref_sheet}"'
                            )

                if is_empty(ref_column):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'ReferenceDataColumn is required for REFERENCE type'
                    )

            # --- DUPLICATE2: ReferenceDataColumn required, ReferenceDataSheet not used ---
            if vtype == "DUPLICATE2":
                ref_column = clean_field(r.get("ReferenceDataColumn"))
                sheet_name = str(r.get("SheetName", "")).strip()

                if is_empty(ref_column):
                    self.errors.append(
                        f'{VALIDATION_RULES_SHEET} row {row}: '
                        f'ReferenceDataColumn is required for DUPLICATE2 type '
                        f'(defines the second column to check)'
                    )
                elif sheet_name and sheet_exists(self.workbook, sheet_name):
                    if find_column_index(get_headers(get_sheet(self.workbook, sheet_name)), ref_column) is None:
                        if not self._is_enrichment_target(sheet_name, ref_column):
                            self.errors.append(
                                f'{VALIDATION_RULES_SHEET} row {row}: '
                                f'ReferenceDataColumn "{ref_column}" not found '
                                f'in sheet "{sheet_name}"'
                            )
                    elif normalize_string(ref_column) == normalize_string(r.get("ColumnName", "")):
                        self.errors.append(
                            f'{VALIDATION_RULES_SHEET} row {row}: '
                            f'ReferenceDataColumn "{ref_column}" is the same as '
                            f'ColumnName for a DUPLICATE2 rule — use ValidationType '
                            f'DUPLICATE instead for a single-column duplicate check'
                        )

        # --- Uniqueness ---
        self._check_unique(seq_nums,  VALIDATION_RULES_SHEET, "SequenceNum")
        self._check_unique(rule_codes, VALIDATION_RULES_SHEET, "RuleCode")
        self._check_unique(rule_names, VALIDATION_RULES_SHEET, "RuleName")

    # -----------------------------------------------------------------------
    # EnrichmentRules
    # -----------------------------------------------------------------------

    def _validate_enrichment_rules(self) -> None:

        if not sheet_exists(self.workbook, ENRICHMENT_RULES_SHEET):
            self.logger.debug(f'Sheet "{ENRICHMENT_RULES_SHEET}" not found — skipping')
            return

        sheet, headers = self._sheet_and_headers(ENRICHMENT_RULES_SHEET)

        header_error_count = 0
        for col in MANDATORY_ENRICHMENT_COLUMNS:
            if find_column_index(headers, col) is None:
                self.errors.append(
                    f'{ENRICHMENT_RULES_SHEET}: mandatory column "{col}" not found'
                )
                header_error_count += 1

        for col in CONDITIONAL_ENRICHMENT_COLUMNS:
            if find_column_index(headers, col) is None:
                self.errors.append(
                    f'{ENRICHMENT_RULES_SHEET}: column "{col}" not found. '
                    f'This column may be left empty on individual rows, but '
                    f'its header must exist in the sheet — otherwise every '
                    f'row silently falls back to a default value, which can '
                    f'hide configuration mistakes such as a renamed column.'
                )
                header_error_count += 1

        if header_error_count:
            return

        seq_nums, rule_codes, rule_names = [], [], []
        all_enrichment_columns = MANDATORY_ENRICHMENT_COLUMNS + CONDITIONAL_ENRICHMENT_COLUMNS

        for row in range(2, last_data_row(sheet) + 1):

            r = self._row_dict(sheet, headers, row, all_enrichment_columns)

            if self._all_empty(r):
                continue

            active = str(r.get("Active", "")).strip().upper()
            if active not in ("YES", "NO"):
                self.errors.append(
                    f'{ENRICHMENT_RULES_SHEET} row {row}: Active must be YES or NO'
                )
                continue

            if active == "NO":
                continue

            # --- Skip rows depending on a SheetName that failed to import (FS 8.5.4) ---
            row_sheet_name = str(r.get("SheetName", "")).strip()
            if normalize_string(row_sheet_name) in self.failed_sheets:
                self.logger.warning(
                    f'{ENRICHMENT_RULES_SHEET} row {row}: row skipped: depends on '
                    f'SheetName "{row_sheet_name}" which failed to import'
                )
                continue

            # --- SequenceNum ---
            if is_empty(r["SequenceNum"]):
                self.errors.append(f'{ENRICHMENT_RULES_SHEET} row {row}: SequenceNum is empty')
            else:
                try:
                    seq_nums.append(int(r["SequenceNum"]))
                except (ValueError, TypeError):
                    self.errors.append(
                        f'{ENRICHMENT_RULES_SHEET} row {row}: SequenceNum must be numeric'
                    )

            self._check_required(r, row, ENRICHMENT_RULES_SHEET, "RuleCode", rule_codes)
            self._check_required(r, row, ENRICHMENT_RULES_SHEET, "RuleName", rule_names)

            # --- SheetName / TargetColumn ---
            sheet_name  = str(r.get("SheetName", "")).strip()
            target_col  = str(r.get("TargetColumn", "")).strip()

            if sheet_name and not sheet_exists(self.workbook, sheet_name):
                self.errors.append(
                    f'{ENRICHMENT_RULES_SHEET} row {row}: '
                    f'SheetName "{sheet_name}" not found in workbook'
                )
            elif sheet_name and target_col:
                # TargetColumn must NOT already exist — enrichment creates new columns only
                existing_headers = get_headers(get_sheet(self.workbook, sheet_name))
                if find_column_index(existing_headers, target_col) is not None:
                    self.errors.append(
                        f'{ENRICHMENT_RULES_SHEET} row {row}: '
                        f'TargetColumn "{target_col}" already exists in sheet "{sheet_name}". '
                        f'Enrichment rules must write to new columns only to prevent overwriting existing data.'
                    )

            # --- Color (optional for EnrichmentRules, FS 8.7.2) ---
            self._check_color(r, row, ENRICHMENT_RULES_SHEET, required=False)

            # --- CustomFunctionName ---
            fn_name = str(r.get("CustomFunctionName", "")).strip()
            if is_empty(fn_name):
                self.errors.append(
                    f'{ENRICHMENT_RULES_SHEET} row {row}: CustomFunctionName is required'
                )
            elif not self.registry.has(fn_name):
                self.errors.append(
                    f'{ENRICHMENT_RULES_SHEET} row {row}: '
                    f'CustomFunctionName "{fn_name}" not found in function registry'
                )
            else:
                # Check REQUIRED_PARAMS against FunctionArguments
                required = self.registry.get_required_params(fn_name)
                if required:
                    fn_args_str = str(r.get("FunctionArguments", "") or "").strip()
                    if not fn_args_str:
                        self.errors.append(
                            f'{ENRICHMENT_RULES_SHEET} row {row}: '
                            f'FunctionArguments is required for function "{fn_name}" '
                            f'(required parameters: {", ".join(required)})'
                        )
                    else:
                        # Parse and check each required param is present
                        provided = set()
                        for part in fn_args_str.split(";"):
                            if "=" in part:
                                provided.add(part.split("=")[0].strip().lower())
                        missing = [p for p in required if p.lower() not in provided]
                        if missing:
                            self.errors.append(
                                f'{ENRICHMENT_RULES_SHEET} row {row}: '
                                f'FunctionArguments for "{fn_name}" is missing '
                                f'required parameter(s): {", ".join(missing)}'
                            )

        self._check_unique(seq_nums,   ENRICHMENT_RULES_SHEET, "SequenceNum")
        self._check_unique(rule_codes, ENRICHMENT_RULES_SHEET, "RuleCode")
        self._check_unique(rule_names, ENRICHMENT_RULES_SHEET, "RuleName")

    # -----------------------------------------------------------------------
    # Helpers
    # -----------------------------------------------------------------------

    def _sheet_and_headers(self, sheet_name: str):
        sheet   = get_sheet(self.workbook, sheet_name)
        headers = get_headers(sheet)
        return sheet, headers

    def _row_dict(self, sheet, headers: dict, row: int, canonical_columns: list) -> dict:
        """
        Builds the row dict keyed by CANONICAL column name (e.g. "SequenceNum"),
        resolving each canonical name to its actual column index case-
        insensitively. This is what makes every downstream r.get("SequenceNum")
        / r["SequenceNum"] work correctly regardless of how the header is
        actually cased in the file (FS section 9: column names are
        case-insensitive). Only call this after confirming all of
        canonical_columns exist via find_column_index — a canonical column
        with no match resolves to None here rather than raising.
        """
        result = {}
        for canon in canonical_columns:
            idx = find_column_index(headers, canon)
            result[canon] = cell_value(sheet.cell(row=row, column=idx)) if idx is not None else None
        return result

    def _all_empty(self, row_data: dict) -> bool:
        return all(is_empty(v) for v in row_data.values())

    def _check_required(
        self, r: dict, row: int, sheet_name: str, col: str, collector: list
    ) -> None:
        value = r.get(col)
        if is_empty(value):
            self.errors.append(f'{sheet_name} row {row}: {col} is empty')
        else:
            collector.append(str(value).strip().lower())

    def _check_color(
        self, r: dict, row: int, sheet_name: str, required: bool
    ) -> None:
        """
        Color must be a known name (built-in or config predefined_colors) or a
        plain 6-digit hex code with no "#" (FS 8.7.2 / 8.7.3). "#RRGGBB" and
        3-digit shorthand are rejected here so the fill-writing step can use
        the value as-is. Empty is an error only when required.
        """
        color = str(r.get("Color", "") or "").strip().upper()
        if not color:
            if required:
                self.errors.append(f'{sheet_name} row {row}: Color is empty')
            return

        if (
            color not in self.config.color_map
            and not re.fullmatch(r'[0-9A-F]{6}', color)
        ):
            self.errors.append(
                f'{sheet_name} row {row}: '
                f'Color "{color}" is not a known color name or a 6-digit hex '
                f'code without "#" (e.g. FF0000)'
            )

    def _check_unique(self, values: list, sheet_name: str, col: str) -> None:
        if len(values) != len(set(values)):
            self.errors.append(
                f'{sheet_name}: duplicate {col} values detected'
            )

    def _is_enrichment_target(self, sheet_name: str, col_name: str) -> bool:
        """
        Returns True if col_name is a TargetColumn of any active EnrichmentRule
        for the given sheet. These columns don't exist yet at validation time
        but will be created during enrichment execution.
        """
        if not sheet_exists(self.workbook, ENRICHMENT_RULES_SHEET):
            return False
        sheet   = get_sheet(self.workbook, ENRICHMENT_RULES_SHEET)
        headers = get_headers(sheet)
        target_col_idx = find_column_index(headers, "TargetColumn")
        sheet_name_idx = find_column_index(headers, "SheetName")
        active_idx     = find_column_index(headers, "Active")
        if target_col_idx is None or sheet_name_idx is None:
            return False
        for row in range(2, last_data_row(sheet) + 1):
            active = str(
                (cell_value(sheet.cell(row=row, column=active_idx)) if active_idx else None) or ""
            ).strip().upper()
            if active != "YES":
                continue
            rule_sheet = str(cell_value(sheet.cell(row=row, column=sheet_name_idx)) or "").strip()
            target_col = str(cell_value(sheet.cell(row=row, column=target_col_idx)) or "").strip()
            if normalize_string(rule_sheet) == normalize_string(sheet_name) and \
               normalize_string(target_col) == normalize_string(col_name):
                return True
        return False
