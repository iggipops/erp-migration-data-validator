import math
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from utils.logger import get_logger


# ---------------------------------------------------------------------------
# Workbook open / copy
# ---------------------------------------------------------------------------

def copy_workbook(source_path: str, output_path: str) -> None:
    """Copy input file to output path, creating parent dirs if needed."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source_path, output_path)


def open_workbook(path: str) -> Workbook:
    """
    Open an xlsx workbook with openpyxl (keep_vba=False) — twice (F-01).

    The returned workbook is the normal one: formulas are preserved, and
    it is the workbook that gets highlighted and saved. A second,
    data_only=True copy (the last-calculated value Excel cached for each
    formula cell) is attached to it as its "values twin"; validators and
    enrichment read business values from the twin via cell_value() /
    sheet_to_dataframe(), never the formula text. Any formula cell whose
    cached value is missing is logged as a warning.
    """
    wb = load_workbook_pair(path)
    warn_missing_cached_values(wb)
    return wb


def load_workbook_pair(path: str) -> Workbook:
    """open_workbook() without the missing-cached-value scan."""
    wb = load_workbook(path)
    attach_values_twin(wb, load_workbook(path, data_only=True))
    return wb


def save_workbook(workbook: Workbook, path: str) -> None:
    """Save the workbook to the given path."""
    workbook.save(path)


# ---------------------------------------------------------------------------
# Values twin (F-01)
#
# The output workbook keeps formulas (it is what gets saved). Business
# values are read from a data_only copy of the same workbook — the
# "twin" — found through the workbook itself, so a plain Workbook with no
# twin attached (tests, generated sheets) reads exactly as before.
# ---------------------------------------------------------------------------

_TWIN_ATTR = "_values_twin"


def attach_values_twin(workbook: Workbook, twin: Workbook) -> None:
    setattr(workbook, _TWIN_ATTR, twin)


def get_values_twin(workbook: Workbook) -> Workbook | None:
    return getattr(workbook, _TWIN_ATTR, None)


def values_sheet(sheet: Worksheet, create: bool = False) -> Worksheet | None:
    """
    The twin's counterpart of `sheet`, or `sheet` itself if its workbook
    has no twin. If a twin exists but lacks the sheet: None, or a new
    empty twin sheet when create=True.
    """
    twin = get_values_twin(sheet.parent)
    if twin is None:
        return sheet
    if sheet.title in twin.sheetnames:
        return twin[sheet.title]
    return twin.create_sheet(sheet.title) if create else None


def cell_value(cell):
    """
    The value a validator should check: for a formula cell, the cached
    (last-calculated) result from the values twin; otherwise the cell's
    own value. A formula whose cached value is missing yields None (the
    warning is logged when the workbook is opened).
    """
    if cell.data_type != "f":
        return cell.value
    twin_sheet = values_sheet(cell.parent)
    if twin_sheet is None:
        return None
    return twin_sheet.cell(row=cell.row, column=cell.column).value


def set_cell_value(sheet: Worksheet, row: int, column: int, value) -> None:
    """Write a plain value to the sheet and its values twin, so both stay in step."""
    sheet.cell(row=row, column=column, value=value)
    twin_sheet = values_sheet(sheet, create=True)
    if twin_sheet is not sheet:
        twin_sheet.cell(row=row, column=column, value=value)


def warn_missing_cached_values(workbook: Workbook, sheet_names: list[str] | None = None,
                               source: str = "") -> int:
    """
    Log one warning per sheet that has formula cells with no cached value
    under data_only=True — the file was never opened and saved in real
    Excel, so its formulas have no calculated results. Returns the total
    number of such cells.
    """
    logger = get_logger()
    total = 0
    for name in (sheet_names if sheet_names is not None else workbook.sheetnames):
        sheet = workbook[name]
        twin_sheet = values_sheet(sheet)
        if twin_sheet is None or twin_sheet is sheet:
            continue
        missing = [
            cell.coordinate
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
            and twin_sheet.cell(row=cell.row, column=cell.column).value is None
        ]
        if missing:
            total += len(missing)
            shown = ", ".join(missing[:5]) + (", ..." if len(missing) > 5 else "")
            logger.warning(
                f"{source}Sheet '{name}': {len(missing)} formula cell(s) have no "
                f"cached value ({shown}) — the file was never opened and saved in "
                f"Excel, so these cells have no calculated result to validate"
            )
    return total


def last_data_row(sheet: Worksheet) -> int:
    """
    1-based number of the last row of `sheet` holding a real value (F-02),
    or 0 for an empty sheet. openpyxl's max_row also counts rows that only
    carry formatting or a deleted value, so it can run far past the data.
    A formula cell counts as a real value. Scanned per sheet, not per column.
    """
    last = 0
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if any(v is not None and str(v).strip() != "" for v in row):
            last = idx
    return last


# ---------------------------------------------------------------------------
# Sheet helpers
# ---------------------------------------------------------------------------

def find_sheet_name(workbook: Workbook, sheet_name: str) -> str | None:
    """
    Case-insensitive worksheet name lookup (FS section 9: worksheet names
    are case-insensitive framework-wide). Returns the sheet's actual name
    as stored in the workbook (original case preserved) if a case-
    insensitive match exists, else None.
    """
    target = str(sheet_name).strip().lower()
    for name in workbook.sheetnames:
        if name.strip().lower() == target:
            return name
    return None


def sheet_exists(workbook: Workbook, sheet_name: str) -> bool:
    return find_sheet_name(workbook, sheet_name) is not None


def get_sheet(workbook: Workbook, sheet_name: str) -> Worksheet:
    """
    Case-insensitive worksheet lookup. Raises KeyError (matching the
    dict-access contract callers previously relied on from workbook[name])
    if no case-insensitive match exists.
    """
    actual = find_sheet_name(workbook, sheet_name)
    if actual is None:
        raise KeyError(f'Worksheet "{sheet_name}" not found')
    return workbook[actual]


def get_or_create_sheet(workbook: Workbook, sheet_name: str) -> Worksheet:
    """Return sheet if it exists (case-insensitive), create it if not."""
    actual = find_sheet_name(workbook, sheet_name)
    if actual is not None:
        return workbook[actual]
    return workbook.create_sheet(sheet_name)


def remove_sheet_if_exists(workbook: Workbook, sheet_name: str) -> None:
    actual = find_sheet_name(workbook, sheet_name)
    if actual is not None:
        workbook.remove(workbook[actual])


# ---------------------------------------------------------------------------
# Header / row helpers
# ---------------------------------------------------------------------------

def get_headers(sheet: Worksheet) -> dict[str, int]:
    """
    Return {column_name: column_index} from the first row.
    Column names are stripped but NOT lowercased — matching is done
    by callers that need case-insensitive behaviour.
    """
    headers: dict[str, int] = {}
    for cell in sheet[1]:
        value = cell_value(cell)
        if value is not None:
            headers[str(value).strip()] = cell.column
    return headers


def get_headers_lower(sheet: Worksheet) -> dict[str, int]:
    """Same as get_headers but keys are lowercased for case-insensitive lookup."""
    return {k.lower(): v for k, v in get_headers(sheet).items()}


def find_column_index(headers: dict[str, int], col_name: str) -> int | None:
    """
    Case-insensitive column name lookup within a headers dict returned
    by get_headers() (FS section 9: column names are case-insensitive
    framework-wide). Returns the column index if a case-insensitive
    match exists, else None.
    """
    target = str(col_name).strip().lower()
    for name, idx in headers.items():
        if name.strip().lower() == target:
            return idx
    return None


def sheet_to_dataframe(sheet: Worksheet) -> pd.DataFrame:
    """
    Read an openpyxl sheet into a pandas DataFrame.
    First row is used as column headers.
    All values are kept as-is (no dtype coercion): formula cells come from
    the values twin (F-01), rows stop at the last real value (F-02), and
    the frame is dtype=object (F-03) so a float stays a float.
    """
    last_row = last_data_row(sheet)
    if last_row == 0:
        return pd.DataFrame()
    rows = list((values_sheet(sheet) or sheet).iter_rows(
        max_row=last_row, values_only=True))

    headers = [str(c).strip() if c is not None else f"_col{i}"
               for i, c in enumerate(rows[0])]
    data = rows[1:]
    return pd.DataFrame(data, columns=headers, dtype=object)


# ---------------------------------------------------------------------------
# Value helpers
# ---------------------------------------------------------------------------

# Leading characters that Excel may interpret as a formula/expression trigger
# (the classic CSV/formula-injection set), guarded against in append_row_safe().
_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@")


def append_row_safe(sheet: Worksheet, values: list) -> None:
    """
    Append a row of values the tool itself is writing (not values read as
    genuine formulas from an input workbook), forcing any string starting
    with =, +, -, or @ to be stored as literal text.

    openpyxl auto-detects a leading "=" and stores the cell as a live
    formula (data_type 'f'); +, -, and @ are not auto-detected by openpyxl
    but are still formula triggers when the file is opened in Excel. A
    string value like "=HYPERLINK(...)" copied verbatim from an external
    CSV must never become a live/interpreted formula in our output —
    explicitly setting data_type to string ('s') prevents that regardless
    of the leading character.
    """
    sheet.append(values)
    row_idx = sheet.max_row
    for col_idx, value in enumerate(values, start=1):
        if isinstance(value, str) and value[:1] in _FORMULA_TRIGGER_CHARS:
            sheet.cell(row=row_idx, column=col_idx).data_type = "s"


def is_empty(value) -> bool:
    """
    True if value is None, a whitespace-only string, or the text "nan"
    (case-insensitive). The "nan" case covers artifacts produced when
    pandas reads empty Excel formula cells (e.g. VLOOKUP results that
    evaluate to nothing) — per FS section 2.5.2, these are treated as
    empty business data, not as the literal text "nan".
    """
    if value is None:
        return True
    s = str(value).strip()
    if s == "":
        return True
    return s.lower() == "nan"


def clean_field(value) -> str:
    """
    Returns '' if is_empty(value) is True (None, blank, or the literal
    text "nan"), otherwise the stripped string form of value.

    Use this in place of str(value).strip() whenever the result will be
    passed to is_empty(): stringifying None first turns it into the
    literal text "None", which is_empty() does not recognize as empty,
    so the empty case silently falls through as if a real value were set.
    """
    return "" if is_empty(value) else str(value).strip()


def normalize_string(value) -> str:
    """Strip and lowercase. Returns '' for None."""
    if value is None:
        return ""
    return str(value).strip().lower()


def is_numeric_null(value) -> bool:
    """
    True if value represents a numeric null:
    empty/None, or a numeric zero in any decimal format
    (0, 0.0, 0,0, 00, 0.00, etc.).
    Used by the NULL validation type.
    """
    if is_empty(value):
        return True
    normalized = str(value).strip().replace(",", ".")
    try:
        parsed = float(normalized)
        if math.isnan(parsed) or math.isinf(parsed):
            return False
        return parsed == 0.0
    except ValueError:
        return False
