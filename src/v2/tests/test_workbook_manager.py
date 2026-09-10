import pytest
from openpyxl import Workbook
from workbook.workbook_manager import WorkbookManager


def make_input_file(tmp_path):
    src = tmp_path / "input.xlsx"
    wb = Workbook()
    wb.active.append(["ItemId"])
    wb.active.append(["A1"])
    wb.save(src)
    return src


def test_copy_and_open_creates_output_file(tmp_path, fake_config):
    src = make_input_file(tmp_path)
    fake_config.input_file = str(src)
    fake_config.output_file = str(tmp_path / "output.xlsx")

    mgr = WorkbookManager(fake_config)
    wb = mgr.copy_and_open()

    assert (tmp_path / "output.xlsx").exists()
    assert wb.active.cell(row=2, column=1).value == "A1"


def test_save_writes_to_output_path(tmp_path, fake_config):
    src = make_input_file(tmp_path)
    fake_config.input_file = str(src)
    fake_config.output_file = str(tmp_path / "output.xlsx")

    mgr = WorkbookManager(fake_config)
    wb = mgr.copy_and_open()
    wb.active.cell(row=3, column=1, value="A2")
    mgr.save()

    from openpyxl import load_workbook
    reloaded = load_workbook(str(tmp_path / "output.xlsx"))
    assert reloaded.active.cell(row=3, column=1).value == "A2"


def test_save_without_open_raises(fake_config):
    mgr = WorkbookManager(fake_config)
    with pytest.raises(RuntimeError):
        mgr.save()


def test_editing_output_does_not_affect_input(tmp_path, fake_config):
    src = make_input_file(tmp_path)
    fake_config.input_file = str(src)
    fake_config.output_file = str(tmp_path / "output.xlsx")

    mgr = WorkbookManager(fake_config)
    wb = mgr.copy_and_open()
    wb.active.cell(row=2, column=1, value="CHANGED")
    mgr.save()

    from openpyxl import load_workbook
    original = load_workbook(str(src))
    assert original.active.cell(row=2, column=1).value == "A1"
