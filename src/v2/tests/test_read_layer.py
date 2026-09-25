"""
Read-layer tests (FS 8.4 / 8.6.1): F-01 cached formula values, F-02 true last
row, F-03 dtype=object — for the primary workbook and for external xlsx files.
"""
import logging
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from importer.importer import Importer
from validators.builtin import empty, numeric
from workbook.excel_utils import (
    cell_value, last_data_row, open_workbook, set_cell_value, sheet_to_dataframe,
)

RED = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")


def save_with_cache(wb, path, cached=None):
    """
    Save `wb` and, if `cached` = {formula_text: cached_value} is given, write
    those cached results into the file the way Excel does. openpyxl itself
    never writes them, so without this the file looks "never opened in Excel".
    """
    wb.save(path)
    if not cached:
        return
    src = path.with_suffix(".src")
    path.rename(src)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                xml = data.decode("utf-8")
                for formula, value in cached.items():
                    xml = xml.replace(f"<f>{formula}</f><v />", f"<f>{formula}</f><v>{value}</v>")
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    src.unlink()


def formula_book(path, cached=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(["Qty", "Double"])
    ws.append([5, "=A2*2"])
    ws.append([7, "=A3*2"])
    save_with_cache(wb, path, {"A2*2": 10, "A3*2": 14} if cached else None)
    return path


# --- F-01: cached values for formulas, formulas kept for output ------------

def test_cell_value_returns_cached_value_and_output_keeps_formula(tmp_path):
    wb = open_workbook(str(formula_book(tmp_path / "in.xlsx")))
    cell = wb["Items"]["B2"]
    assert cell.value == "=A2*2"          # output side keeps the formula
    assert cell_value(cell) == 10         # validators see the calculated value

    out = tmp_path / "out.xlsx"
    wb.save(out)
    assert load_workbook(out)["Items"]["B2"].value == "=A2*2"


def test_missing_cached_value_logs_warning_and_reads_as_none(tmp_path, caplog):
    with caplog.at_level(logging.WARNING, logger="erp_migration_data_validator"):
        wb = open_workbook(str(formula_book(tmp_path / "in.xlsx", cached=False)))
    assert any("no cached value" in r.message and "Items" in r.message for r in caplog.records)
    assert cell_value(wb["Items"]["B2"]) is None


def test_no_warning_when_cached_values_present(tmp_path, caplog):
    with caplog.at_level(logging.WARNING, logger="erp_migration_data_validator"):
        open_workbook(str(formula_book(tmp_path / "in.xlsx")))
    assert not any("no cached value" in r.message for r in caplog.records)


def test_numeric_validation_checks_cached_value_not_formula_text(tmp_path, fake_config, issues_writer_factory):
    wb = open_workbook(str(formula_book(tmp_path / "in.xlsx")))
    iw = issues_writer_factory(wb, fake_config)
    rule = {"RuleCode": "R1", "RuleName": "n", "Severity": "ERROR", "Color": "RED"}
    numeric.validate(wb["Items"], 2, "Double", rule, {"issues_writer": iw, "config": fake_config})
    assert iw.issue_counter == 0          # would be 2 ("=A2*2" is text) without the twin


def test_issue_highlights_output_cell_and_reports_cached_value(tmp_path, fake_config, issues_writer_factory):
    path = tmp_path / "in.xlsx"
    wb = Workbook()
    wb.active.title = "Items"
    wb.active.append(["Id", "Calc"])
    wb.active.append(["A1", '=IF(1=1,"","x")'])
    save_with_cache(wb, path, {'IF(1=1,"","x")': ""})
    wb = open_workbook(str(path))
    iw = issues_writer_factory(wb, fake_config)
    rule = {"RuleCode": "R1", "RuleName": "e", "Severity": "ERROR", "Color": "RED"}
    empty.validate(wb["Items"], 2, "Calc", rule, {"issues_writer": iw})
    assert iw.issue_counter == 1
    assert wb["Items"]["B2"].fill.start_color.rgb.endswith("FF0000")


def test_sheet_to_dataframe_reads_cached_values(tmp_path):
    wb = open_workbook(str(formula_book(tmp_path / "in.xlsx")))
    assert sheet_to_dataframe(wb["Items"])["Double"].tolist() == [10, 14]


def test_set_cell_value_keeps_twin_in_step(tmp_path):
    wb = open_workbook(str(formula_book(tmp_path / "in.xlsx")))
    set_cell_value(wb["Items"], 1, 3, "Enriched")
    set_cell_value(wb["Items"], 2, 3, "x")
    df = sheet_to_dataframe(wb["Items"])
    assert list(df.columns) == ["Qty", "Double", "Enriched"]
    assert df["Enriched"].tolist()[0] == "x"


# --- F-01 against a real Excel-saved file (not a hand-crafted fixture) -----
#
# Every case above builds its workbook with openpyxl and patches cached
# values into the raw XML via save_with_cache() — useful for control over
# exactly which cells have/lack a cached value, but it says nothing about
# how a real spreadsheet application actually serializes one. This fixture
# was opened and saved by Microsoft Excel (Online) itself, so its cached
# values are genuinely Excel's.
#
# A sibling file, "Stock Balanse Check.xlsx" (no "2"), is the same
# spreadsheet saved by LibreOffice instead of Excel — both are kept in
# test_data/v2/ deliberately, in case the two applications' formula-caching
# behavior ever needs comparing again.
#
# Neither file is in test_data/v2/demo/, so both stay private-repo-only —
# this test is expected to skip, not fail, in the public repo.

REPO_ROOT = Path(__file__).resolve().parents[3]
REAL_EXCEL_WORKBOOK = REPO_ROOT / "test_data" / "v2" / "Stock Balanse Check2.xlsx"


@pytest.mark.skipif(
    not REAL_EXCEL_WORKBOOK.exists(),
    reason="real-Excel fixture (Stock Balanse Check2.xlsx) not present in "
           "this repo — private-repo-only, see test docstring.",
)
def test_real_excel_file_cached_values_resolve_via_cell_value(caplog):
    """
    Runs a genuinely Excel-saved workbook through the real open_workbook()/
    cell_value() code path (not plain openpyxl, not save_with_cache()'s
    synthetic XML patcher). A representative subset of formula cells is
    checked, not exhaustive coverage — the point is exercising the real
    code path against a real Excel file, not covering every cell.
    """
    with caplog.at_level(logging.WARNING, logger="erp_migration_data_validator"):
        wb = open_workbook(str(REAL_EXCEL_WORKBOOK))

    data = wb["Data"]
    wmsl = wb["WMSL"]

    # Cached values, as Excel computed them — read via cell_value().
    assert cell_value(data["R2"]) == "БезНомеров"
    assert cell_value(data["S2"]) == "2-7-1-3"
    assert cell_value(data["T2"]) == "Сборка печатной платы"
    assert cell_value(wmsl["A2"]) == "WHS1-1-1-2"
    assert cell_value(wmsl["A3"]) == "WHS1-1-3-2"

    # Formulas themselves are preserved on the output side (F-01's other
    # half): the cell's own .value is still the formula, not the cached
    # result — that's only reachable through cell_value().
    assert data["R2"].value == "=VLOOKUP(J2,Items!$A:$G,7,0)"
    assert wmsl["A2"].value == "=B2&C2"

    # None of this file's formula cells should have hit the missing-
    # cached-value warning path — Excel cached every one of them.
    assert not any("no cached value" in r.message for r in caplog.records)


# --- F-02: true last row ----------------------------------------------------

def test_last_data_row_ignores_formatting_only_rows():
    wb = Workbook()
    ws = wb.active
    ws.append(["Id"])
    ws.append(["A1"])
    ws.append(["A2"])
    ws["A50"].fill = RED                   # formatting only: max_row becomes 50
    assert ws.max_row == 50
    assert last_data_row(ws) == 3


def test_last_data_row_edge_cases():
    wb = Workbook()
    ws = wb.active
    assert last_data_row(ws) == 0                       # empty sheet
    ws["A1"] = "h"
    ws["A5"] = "   "                                    # whitespace is not a value
    assert last_data_row(ws) == 1
    ws["B4"] = 0                                        # zero is a value; found in another column
    assert last_data_row(ws) == 4
    ws["A6"] = "=1+1"                                   # a formula counts
    assert last_data_row(ws) == 6


def test_validators_skip_phantom_trailing_rows(fake_config, issues_writer_factory):
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(["Id", "Other"])
    ws.append(["A1", "x"])
    ws.append([None, "x"])                              # one real empty cell
    ws["B60"].fill = RED                                # 57 phantom rows below
    iw = issues_writer_factory(wb, fake_config)
    rule = {"RuleCode": "R1", "RuleName": "e", "Severity": "ERROR", "Color": "RED"}
    empty.validate(ws, 1, "Id", rule, {"issues_writer": iw})
    assert iw.issue_counter == 1


def test_sheet_to_dataframe_stops_at_last_real_row():
    wb = Workbook()
    ws = wb.active
    ws.append(["Id"])
    ws.append(["A1"])
    ws["A40"].fill = RED
    assert len(sheet_to_dataframe(ws)) == 1


# --- F-03: dtype=object -----------------------------------------------------

def test_sheet_to_dataframe_keeps_ints_and_floats_as_read():
    wb = Workbook()
    ws = wb.active
    ws.append(["Qty", "Price", "Note"])
    ws.append([5, 1.5, "a"])
    ws.append([None, 2.0, "b"])
    df = sheet_to_dataframe(ws)
    assert all(t == object for t in df.dtypes)
    assert df["Qty"].tolist() == [5, None]              # not [5.0, nan]
    assert type(df["Qty"][0]) is int
    assert df["Price"].tolist() == [1.5, 2.0]
    assert type(df["Price"][1]) is float                # 2.0 stays exactly as read


# --- External xlsx goes through the same read logic -------------------------

def ext_rule(path, sheet="Ext", original="Data"):
    return {"SequenceNum": 1, "SheetName": sheet, "FilePath": str(path),
            "Format": "xlsx", "Active": "Yes", "OriginalSheetName": original,
            "CSVDelimiter": "", "CSVEncoding": ""}


def external_book(path, cached=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append([" Id ", "Qty", "When", "Calc", "Text"])
    ws.append(["A1", 5, datetime(2026, 1, 31), "=B2*2", "=not a formula"])
    ws.append(["A2", 2.5, None, "=B3*2", "y"])
    ws["A30"].fill = RED
    ws["E2"].data_type = "s"                 # text that merely starts with "="
    save_with_cache(wb, path, {"B2*2": 10, "B3*2": 5} if cached else None)
    return path


def test_external_xlsx_keeps_native_types_and_formulas(make_workbook, fake_config, tmp_path):
    wb = make_workbook()
    result = Importer(wb, fake_config).import_external_files(
        [ext_rule(external_book(tmp_path / "ext.xlsx"))], [])
    assert result["loaded"] == ["Ext"]
    ws = wb["Ext"]
    assert ws["A1"].value == "Id"                                       # header stripped
    assert ws["B2"].value == 5 and isinstance(ws["B2"].value, int)      # not the text "5"
    assert ws["B3"].value == 2.5
    assert ws["C2"].value == datetime(2026, 1, 31)
    assert ws["D2"].value == "=B2*2" and ws["D2"].data_type == "f"      # formula preserved
    assert ws["E2"].data_type == "s"                                    # text stays text
    assert ws.max_row == 3                                              # no phantom rows


def test_external_xlsx_formulas_read_as_cached_values(fake_config, tmp_path):
    wb = Workbook()
    wb.active.title = "Main"
    wb.active.append(["x"])
    main = tmp_path / "main.xlsx"
    wb.save(main)
    wb = open_workbook(str(main))
    Importer(wb, fake_config).import_external_files(
        [ext_rule(external_book(tmp_path / "ext.xlsx"))], [])
    assert cell_value(wb["Ext"]["D2"]) == 10
    df = sheet_to_dataframe(wb["Ext"])
    assert df["Calc"].tolist() == [10, 5]
    assert len(df) == 2


def test_external_xlsx_missing_cached_value_logs_warning(fake_config, tmp_path, caplog):
    wb = Workbook()
    wb.active.title = "Main"
    main = tmp_path / "main.xlsx"
    wb.save(main)
    wb = open_workbook(str(main))
    with caplog.at_level(logging.WARNING, logger="erp_migration_data_validator"):
        Importer(wb, fake_config).import_external_files(
            [ext_rule(external_book(tmp_path / "ext.xlsx", cached=False))], [])
    assert any("no cached value" in r.message and "ext.xlsx" in r.message for r in caplog.records)


def test_external_xlsx_is_not_touched_by_import_spec(make_workbook, fake_config, tmp_path):
    spec = [{"ExternalFileSheetName": "Ext", "ColumnName": "Qty", "TargetType": "text",
             "Active": "Yes"},
            {"ExternalFileSheetName": "Ext", "ColumnName": "Qty", "TargetType": "integer",
             "Active": "Yes"}]
    wb = make_workbook()
    result = Importer(wb, fake_config).import_external_files(
        [ext_rule(external_book(tmp_path / "ext.xlsx"))], spec)
    assert result["conversion_warnings"]["Ext"] == 0
    assert wb["Ext"]["B3"].value == 2.5                 # a real float stays exactly as read


def test_external_xlsx_unknown_original_sheet_is_a_failed_import(make_workbook, fake_config, tmp_path):
    wb = make_workbook()
    result = Importer(wb, fake_config).import_external_files(
        [ext_rule(external_book(tmp_path / "ext.xlsx"), original="Nope")], [])
    assert result["loaded"] == []
    assert result["failed_sheets"] == {"Ext"}


def test_external_csv_import_unchanged_and_mirrored_to_twin(fake_config, tmp_path):
    main = tmp_path / "main.xlsx"
    wb = Workbook()
    wb.active.title = "Main"
    wb.save(main)
    wb = open_workbook(str(main))
    f = tmp_path / "items.csv"
    f.write_text("ItemId,Qty\nA1,5\n")
    rule = {"SequenceNum": 1, "SheetName": "Items", "FilePath": str(f), "Format": "csv",
            "Active": "Yes", "CSVDelimiter": ",", "CSVEncoding": "utf-8", "OriginalSheetName": ""}
    Importer(wb, fake_config).import_external_files([rule], [])
    assert wb["Items"]["B2"].value == "5"               # csv stays text
    assert sheet_to_dataframe(wb["Items"])["Qty"].tolist() == ["5"]
