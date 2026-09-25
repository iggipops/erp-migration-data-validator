import pytest
from openpyxl import load_workbook

from output.issues_writer import IssuesWriter, ISSUES_HEADERS


def test_initializes_issues_sheet(make_workbook, fake_config):
    wb = make_workbook()
    iw = IssuesWriter(wb, fake_config)
    ws = wb["Issues"]
    assert [c.value for c in ws[1]] == ISSUES_HEADERS


def test_replaces_existing_issues_sheet(make_workbook, fake_config):
    wb = make_workbook()
    wb.create_sheet("Issues").append(["Stale"])
    IssuesWriter(wb, fake_config)
    ws = wb["Issues"]
    assert ws.cell(row=1, column=1).value == "IssueNum"


def test_record_appends_row_and_increments_counters(make_workbook, fake_config):
    wb = make_workbook()
    sheet = wb.create_sheet("Items")
    sheet.append(["ItemId"])
    sheet.append(["A1"])
    iw = IssuesWriter(wb, fake_config)
    rule = {"RuleCode": "R1", "RuleName": "Rule", "Severity": "ERROR", "Color": "RED"}
    cell = sheet.cell(row=2, column=1)
    iw.record(cell=cell, rule=rule, sheet_name="Items", column_name="ItemId",
               row_number=2, cell_content="A1")
    assert iw.issue_counter == 1
    assert iw.error_counter == 1
    assert iw.warning_counter == 0
    issues_ws = wb["Issues"]
    assert issues_ws.cell(row=2, column=1).value == 1
    assert issues_ws.cell(row=2, column=6).value == "R1"


def test_warning_severity_increments_warning_counter(make_workbook, fake_config):
    wb = make_workbook()
    sheet = wb.create_sheet("Items")
    sheet.append(["ItemId"])
    sheet.append(["A1"])
    iw = IssuesWriter(wb, fake_config)
    rule = {"RuleCode": "R1", "RuleName": "Rule", "Severity": "WARNING", "Color": "YELLOW"}
    cell = sheet.cell(row=2, column=1)
    iw.record(cell=cell, rule=rule, sheet_name="Items", column_name="ItemId",
               row_number=2, cell_content="A1")
    assert iw.warning_counter == 1
    assert iw.error_counter == 0


def test_record_leading_equals_cell_content_not_written_as_formula(make_workbook, fake_config):
    """cell_content originating as flagged text (e.g. '=HYPERLINK(...)' from
    an external CSV) must be stored as literal text in the Issues sheet, not
    a live formula (formula/CSV injection guard)."""
    wb = make_workbook()
    sheet = wb.create_sheet("Items")
    sheet.append(["Note"])
    sheet.append(['=HYPERLINK("http://evil")'])
    iw = IssuesWriter(wb, fake_config)
    rule = {"RuleCode": "R1", "RuleName": "Rule", "Severity": "ERROR", "Color": "RED"}
    cell = sheet.cell(row=2, column=1)
    iw.record(cell=cell, rule=rule, sheet_name="Items", column_name="Note",
               row_number=2, cell_content='=HYPERLINK("http://evil")')
    content_cell = wb["Issues"].cell(row=2, column=5)
    assert content_cell.data_type == "s"
    assert content_cell.value == '=HYPERLINK("http://evil")'


def test_record_highlights_cell(make_workbook, fake_config):
    wb = make_workbook()
    sheet = wb.create_sheet("Items")
    sheet.append(["ItemId"])
    sheet.append(["A1"])
    iw = IssuesWriter(wb, fake_config)
    rule = {"RuleCode": "R1", "RuleName": "Rule", "Severity": "ERROR", "Color": "RED"}
    cell = sheet.cell(row=2, column=1)
    iw.record(cell=cell, rule=rule, sheet_name="Items", column_name="ItemId",
               row_number=2, cell_content="A1")
    assert cell.fill.start_color.rgb.endswith("FF0000")


@pytest.mark.parametrize("color, expected_rgb", [
    ("RED",    "FF0000"),   # built-in name
    ("red",    "FF0000"),   # names are case-insensitive
    ("PURPLE", "7030A0"),   # config predefined_colors name
    ("00FF00", "00FF00"),   # plain 6-digit hex
    ("00ff00", "00FF00"),   # lowercase hex
])
def test_record_fill_for_every_color_form_validation_allows(
    make_workbook, fake_config, tmp_path, color, expected_rgb
):
    """FS 8.7.3 guarantees ValidationRules.Color is a known name or a plain
    6-digit hex by the time a rule reaches the IssuesWriter. _get_fill must
    turn each of those forms into a solid fill of the right color that
    survives a save/load round-trip."""
    fake_config.color_map = {**fake_config.color_map, "PURPLE": "7030A0"}
    wb = make_workbook()
    sheet = wb.create_sheet("Items")
    sheet.append(["ItemId"])
    sheet.append(["A1"])
    iw = IssuesWriter(wb, fake_config)
    rule = {"RuleCode": "R1", "RuleName": "Rule", "Severity": "ERROR", "Color": color}
    iw.record(cell=sheet.cell(row=2, column=1), rule=rule, sheet_name="Items",
              column_name="ItemId", row_number=2, cell_content="A1")

    path = tmp_path / "out.xlsx"
    wb.save(path)
    fill = load_workbook(path)["Items"].cell(row=2, column=1).fill
    assert fill.fill_type == "solid"
    assert fill.start_color.rgb == "00" + expected_rgb
    assert fill.end_color.rgb == "00" + expected_rgb
