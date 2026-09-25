import pytest
from openpyxl import load_workbook

from engine.enrichment_engine import EnrichmentEngine


def make_rule(**overrides):
    rule = {
        "SequenceNum": 1, "RuleCode": "E1", "RuleName": "Enrich",
        "SheetName": "Items", "TargetColumn": "Seq",
        "CustomFunctionName": "counter", "FunctionArguments": "",
        "Active": "Yes", "Color": "",
    }
    rule.update(overrides)
    return rule


def test_creates_new_column_with_function_output(make_workbook, sheet_factory, fake_config, fn_registry):
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"], ["A2"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    stats = engine.execute([make_rule()])
    assert stats[0]["status"] == "executed"
    ws = wb["Items"]
    assert ws.cell(row=1, column=2).value == "Seq"
    assert ws.cell(row=2, column=2).value == 1
    assert ws.cell(row=3, column=2).value == 2


def test_new_column_placed_after_last_data_column_not_formatting(
    make_workbook, sheet_factory, fake_config, fn_registry
):
    # F-02 for columns: a formatting-only cell far to the right inflates
    # max_column; the new column must still land right after the data.
    from openpyxl.styles import PatternFill
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"], ["A2"]])
    ws = wb["Items"]
    ws["J1"].fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    assert ws.max_column == 10
    EnrichmentEngine(wb, fake_config, fn_registry).execute([make_rule()])
    assert ws.cell(row=1, column=2).value == "Seq"
    assert ws.cell(row=3, column=2).value == 2
    assert ws.cell(row=1, column=11).value is None


def test_inactive_rule_skipped(make_workbook, sheet_factory, fake_config, fn_registry):
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    stats = engine.execute([make_rule(Active="No")])
    assert stats == []


def test_sheet_resolved_case_insensitively(make_workbook, sheet_factory, fake_config, fn_registry):
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    stats = engine.execute([make_rule(SheetName="items")])
    assert stats[0]["status"] == "executed"
    assert wb["Items"].cell(row=1, column=2).value == "Seq"


def test_unknown_function_marks_failed(make_workbook, sheet_factory, fake_config, fn_registry):
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    stats = engine.execute([make_rule(CustomFunctionName="no_such_function")])
    assert stats[0]["status"] == "failed"


def test_rule_depending_on_failed_import_is_skipped_not_failed(make_workbook, sheet_factory, fake_config, fn_registry):
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    stats = engine.execute([make_rule(SheetName="Stock")], failed_sheets={"Stock"})
    assert stats[0]["status"] == "skipped"


def test_unrelated_enrichment_rule_still_executes_when_other_sheet_failed(
    make_workbook, sheet_factory, fake_config, fn_registry
):
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    rules = [
        make_rule(RuleCode="E1", SheetName="Stock"),
        make_rule(RuleCode="E2", SheetName="Items"),
    ]
    stats = engine.execute(rules, failed_sheets={"Stock"})
    by_code = {s["RuleCode"]: s for s in stats}
    assert by_code["E1"]["status"] == "skipped"
    assert by_code["E2"]["status"] == "executed"
    assert wb["Items"].cell(row=1, column=2).value == "Seq"


def test_function_arguments_string_parsed(make_workbook, sheet_factory, fake_config, fn_registry):
    wb = make_workbook()
    sheet_factory(wb, "Items", [["First", "Last"], ["A", "Smith"], ["B", "Jones"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    rule = make_rule(
        TargetColumn="FullName",
        CustomFunctionName="concatenate",
        FunctionArguments="columns=First,Last;separator=_",
    )
    stats = engine.execute([rule])
    assert stats[0]["status"] == "executed"
    ws = wb["Items"]
    assert ws.cell(row=2, column=3).value == "A_Smith"
    assert ws.cell(row=3, column=3).value == "B_Jones"


def test_function_arguments_quoted_space_separator_preserved(make_workbook, sheet_factory, fake_config, fn_registry):
    # An unquoted separator=" " is indistinguishable from separator= once the
    # cell/args string is trimmed — quoting is required to keep the space.
    wb = make_workbook()
    sheet_factory(wb, "Items", [["First", "Last"], ["A", "Smith"], ["B", "Jones"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    rule = make_rule(
        TargetColumn="FullName",
        CustomFunctionName="concatenate",
        FunctionArguments='columns=First,Last;separator=" "',
    )
    stats = engine.execute([rule])
    assert stats[0]["status"] == "executed"
    ws = wb["Items"]
    assert ws.cell(row=2, column=3).value == "A Smith"
    assert ws.cell(row=3, column=3).value == "B Jones"


def _saved_column_fills(wb, tmp_path, col_idx, rows):
    path = tmp_path / "out.xlsx"
    wb.save(path)
    ws = load_workbook(path)["Items"]
    return [ws.cell(row=r, column=col_idx).fill for r in range(1, rows + 1)]


@pytest.mark.parametrize("color", ["", None, "   "])
def test_blank_color_applies_no_fill_to_generated_column(
    make_workbook, sheet_factory, fake_config, fn_registry, tmp_path, color
):
    """FS 8.8.1: blank EnrichmentRules.Color -> no fill on header or data cells."""
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"], ["A2"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    stats = engine.execute([make_rule(Color=color)])
    assert stats[0]["status"] == "executed"
    for fill in _saved_column_fills(wb, tmp_path, col_idx=2, rows=3):
        assert fill.fill_type is None


@pytest.mark.parametrize("color, expected_rgb", [
    ("RED",    "FF0000"),
    ("PURPLE", "7030A0"),
    ("00FF00", "00FF00"),
])
def test_color_fills_generated_header_and_data_cells(
    make_workbook, sheet_factory, fake_config, fn_registry, tmp_path, color, expected_rgb
):
    """FS 8.8.1: non-blank EnrichmentRules.Color fills header and data cells,
    resolved the same way as ValidationRules.Color."""
    fake_config.color_map = {**fake_config.color_map, "PURPLE": "7030A0"}
    wb = make_workbook()
    sheet_factory(wb, "Items", [["ItemId"], ["A1"], ["A2"]])
    engine = EnrichmentEngine(wb, fake_config, fn_registry)
    stats = engine.execute([make_rule(Color=color)])
    assert stats[0]["status"] == "executed"
    for fill in _saved_column_fills(wb, tmp_path, col_idx=2, rows=3):
        assert fill.fill_type == "solid"
        assert fill.start_color.rgb == "00" + expected_rgb
