import sys
import yaml
from openpyxl import Workbook, load_workbook

import erp_migration_data_validator
import utils.ai_client as ai_client


def build_input_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)

    items = wb.create_sheet("Items")
    items.append(["ItemId", "Qty"])
    items.append(["A1", 5])
    items.append(["A1", 5])   # duplicate ItemId -> should trigger DUPLICATE
    items.append([None, 3])   # empty ItemId -> should trigger EMPTY

    vr = wb.create_sheet("ValidationRules")
    vr.append(["SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName",
               "ValidationType", "Severity", "Color", "Active",
               "ReferenceDataSheet", "ReferenceDataColumn", "CustomFunctionName"])
    vr.append([1, "R1", "No duplicates", "Items", "ItemId", "DUPLICATE", "ERROR", "RED", "Yes", "", "", ""])
    vr.append([2, "R2", "Not empty", "Items", "ItemId", "EMPTY", "ERROR", "RED", "Yes", "", "", ""])

    er = wb.create_sheet("EnrichmentRules")
    er.append(["SequenceNum", "RuleCode", "RuleName", "SheetName", "TargetColumn",
               "CustomFunctionName", "FunctionArguments", "Active", "Color"])
    er.append([1, "E1", "Row number", "Items", "RowNum", "counter", "", "Yes", ""])

    wb.save(path)


def test_full_pipeline_end_to_end(tmp_path, monkeypatch):
    input_file  = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    log_file    = tmp_path / "run.log"
    custom_dir  = tmp_path / "custom_functions"
    custom_dir.mkdir()

    build_input_workbook(input_file)

    config_data = {
        "input_file": str(input_file),
        "output_file": str(output_file),
        "log_file": str(log_file),
        "custom_functions_directory": str(custom_dir),
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config_data))

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["erp_migration_data_validator.py", "--config", str(config_path)])

    exit_code = erp_migration_data_validator.main()

    assert exit_code == 0
    assert output_file.exists()

    result_wb = load_workbook(str(output_file))
    assert "Issues" in result_wb.sheetnames
    assert "Summary" in result_wb.sheetnames

    # Enrichment created the new column
    items = result_wb["Items"]
    header_row = [c.value for c in items[1]]
    assert "RowNum" in header_row

    # Both duplicate ItemId rows AND the empty ItemId row should be flagged
    issues = result_wb["Issues"]
    rule_codes = [
        issues.cell(row=r, column=6).value
        for r in range(2, issues.max_row + 1)
    ]
    assert rule_codes.count("R1") == 2   # both A1 rows
    assert rule_codes.count("R2") == 1   # the empty-ItemId row


def test_pipeline_fails_cleanly_on_bad_config(tmp_path, monkeypatch):
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump({"input_file": "", "output_file": "", "log_file": "", "custom_functions_directory": ""}))

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["erp_migration_data_validator.py", "--config", str(config_path)])

    exit_code = erp_migration_data_validator.main()
    assert exit_code == 1


def build_input_workbook_with_failed_import(path, external_csv_path):
    """
    Workbook with one ExternalFiles rule whose target SheetName ("Bad/Sheet")
    contains a character openpyxl rejects as a sheet title — this makes the
    file itself load fine (so FS 8.2 preflight passes) but fails when the
    Importer tries to create the destination worksheet at Step 5, which is
    a realistic import-time-only failure (preflight doesn't validate
    SheetName's Excel-safety, only emptiness/uniqueness/collision).

    One ValidationRules row targets the sheet that will fail to import;
    an unrelated row targets "Items", which must still run normally.
    """
    wb = Workbook()
    wb.remove(wb.active)

    items = wb.create_sheet("Items")
    items.append(["ItemId", "Qty"])
    items.append(["A1", 5])
    items.append([None, 3])  # empty ItemId -> should trigger EMPTY
    # (a second non-empty column keeps this row from being dropped as a
    # blank trailing row when openpyxl saves/reloads the workbook)

    ef = wb.create_sheet("ExternalFiles")
    ef.append(["SequenceNum", "SheetName", "FilePath", "Format", "Active",
               "OriginalSheetName", "CSVDelimiter"])
    ef.append([1, "Bad/Sheet", str(external_csv_path), "csv", "Yes", "", ","])

    vr = wb.create_sheet("ValidationRules")
    vr.append(["SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName",
               "ValidationType", "Severity", "Color", "Active",
               "ReferenceDataSheet", "ReferenceDataColumn", "CustomFunctionName"])
    vr.append([1, "R1", "Not empty", "Items", "ItemId", "EMPTY", "ERROR", "RED", "Yes", "", "", ""])
    vr.append([2, "R2", "Stock check", "Bad/Sheet", "Qty", "EMPTY", "ERROR", "RED", "Yes", "", "", ""])

    er = wb.create_sheet("EnrichmentRules")
    er.append(["SequenceNum", "RuleCode", "RuleName", "SheetName", "TargetColumn",
               "CustomFunctionName", "FunctionArguments", "Active", "Color"])

    wb.save(path)


def test_one_failed_import_does_not_kill_the_whole_run(tmp_path, monkeypatch):
    """Regression for FS 8.5.4: one bad external file import used to make
    MetadataValidator raise a hard error for every rule referencing that
    SheetName, terminating the ENTIRE run — including unrelated rules on
    other sheets. Now only the dependent rule is skipped."""
    input_file  = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    log_file    = tmp_path / "run.log"
    custom_dir  = tmp_path / "custom_functions"
    custom_dir.mkdir()

    external_csv = tmp_path / "stock.csv"
    external_csv.write_text("Qty\n5\n")

    build_input_workbook_with_failed_import(input_file, external_csv)

    config_data = {
        "input_file": str(input_file),
        "output_file": str(output_file),
        "log_file": str(log_file),
        "custom_functions_directory": str(custom_dir),
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config_data))

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["erp_migration_data_validator.py", "--config", str(config_path)])

    exit_code = erp_migration_data_validator.main()

    # The whole run must NOT be killed by one bad import.
    assert exit_code == 0
    assert output_file.exists()

    result_wb = load_workbook(str(output_file))

    # The unrelated rule on "Items" still ran and produced the correct result.
    issues = result_wb["Issues"]
    rule_codes = [
        issues.cell(row=r, column=6).value
        for r in range(2, issues.max_row + 1)
    ]
    assert rule_codes.count("R1") == 1
    assert "R2" not in rule_codes  # R2 depends on the failed sheet — skipped, never dispatched

    # Summary worksheet distinguishes the failed import and the skipped
    # dependent rule from a genuine runtime failure.
    summary_values = [
        c.value for row in result_wb["Summary"].iter_rows() for c in row if c.value is not None
    ]
    assert any("failed" in str(v) for v in summary_values)
    assert any("skipped" in str(v) for v in summary_values)


def build_input_workbook_with_reference_and_sheetname_dependency(path, external_csv_path):
    """
    Same failed-import trick as build_input_workbook_with_failed_import, but
    exercises BOTH ways a ValidationRules row can depend on a failed sheet:

      R1 — unrelated, on the healthy "Items" sheet — must execute normally.
      R2 — depends on the failed sheet via SheetName — must be skipped.
      R3 — SheetName is the healthy "Items" sheet, but it's a REFERENCE rule
           whose ReferenceDataSheet is the failed sheet — must also be
           skipped, not raise a metadata error and not fall through to a
           runtime failure at Step 8.
    """
    wb = Workbook()
    wb.remove(wb.active)

    items = wb.create_sheet("Items")
    items.append(["ItemId", "Qty"])
    items.append(["A1", 5])
    items.append([None, 3])  # empty ItemId -> should trigger EMPTY (R1)

    ef = wb.create_sheet("ExternalFiles")
    ef.append(["SequenceNum", "SheetName", "FilePath", "Format", "Active",
               "OriginalSheetName", "CSVDelimiter"])
    ef.append([1, "Bad/Sheet", str(external_csv_path), "csv", "Yes", "", ","])

    vr = wb.create_sheet("ValidationRules")
    vr.append(["SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName",
               "ValidationType", "Severity", "Color", "Active",
               "ReferenceDataSheet", "ReferenceDataColumn", "CustomFunctionName"])
    vr.append([1, "R1", "Not empty", "Items", "ItemId", "EMPTY", "ERROR", "RED", "Yes", "", "", ""])
    vr.append([2, "R2", "Stock check", "Bad/Sheet", "Qty", "EMPTY", "ERROR", "RED", "Yes", "", "", ""])
    vr.append([3, "R3", "Item in stock", "Items", "ItemId", "REFERENCE", "ERROR", "RED", "Yes",
               "Bad/Sheet", "Qty", ""])

    er = wb.create_sheet("EnrichmentRules")
    er.append(["SequenceNum", "RuleCode", "RuleName", "SheetName", "TargetColumn",
               "CustomFunctionName", "FunctionArguments", "Active", "Color"])

    wb.save(path)


def test_reference_data_sheet_dependency_isolated_same_as_sheetname_dependency(tmp_path, monkeypatch):
    """
    Follow-up to FS 8.5.4 / test_one_failed_import_does_not_kill_the_whole_run:
    a ValidationRules row can depend on a failed import via SheetName OR via
    ReferenceDataSheet (REFERENCE type). Both must be isolated the same way —
    Step 6 must pass cleanly (zero metadata errors) and Step 8 must skip both
    rows individually — so the whole run still completes normally and the
    unrelated rule produces its normal, correct result.
    """
    input_file  = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    log_file    = tmp_path / "run.log"
    custom_dir  = tmp_path / "custom_functions"
    custom_dir.mkdir()

    external_csv = tmp_path / "stock.csv"
    external_csv.write_text("Qty\n5\n")

    build_input_workbook_with_reference_and_sheetname_dependency(input_file, external_csv)

    config_data = {
        "input_file": str(input_file),
        "output_file": str(output_file),
        "log_file": str(log_file),
        "custom_functions_directory": str(custom_dir),
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config_data))

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["erp_migration_data_validator.py", "--config", str(config_path)])

    exit_code = erp_migration_data_validator.main()

    # The run must complete normally — NOT terminated at Step 6.
    assert exit_code == 0
    assert output_file.exists()

    result_wb = load_workbook(str(output_file))

    # R1 (unrelated, on the healthy "Items" sheet) produced its real result;
    # neither dependent rule was ever dispatched.
    issues = result_wb["Issues"]
    rule_codes = [
        issues.cell(row=r, column=6).value
        for r in range(2, issues.max_row + 1)
    ]
    assert rule_codes.count("R1") == 1
    assert "R2" not in rule_codes
    assert "R3" not in rule_codes

    # Summary's Validation Rules Statistics table (written via the normal
    # write() path, since the run completed) shows the correct status per
    # rule: both dependent rules "skipped", the unrelated rule "executed".
    ws = result_wb["Summary"]

    def status_for(rule_code):
        for row in ws.iter_rows():
            if row[1].value == rule_code:
                return row[5].value
        return None

    assert status_for("R1") == "executed"
    assert status_for("R2") == "skipped"
    assert status_for("R3") == "skipped"


def test_step6_metadata_failure_still_writes_summary_and_saves_output(tmp_path, monkeypatch):
    """FS 8.6: on a Step 6 metadata error, the output workbook (already
    copied at Step 3) must be saved with a Summary worksheet explaining
    the failure — not left as an unlabeled copy."""
    input_file  = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    log_file    = tmp_path / "run.log"
    custom_dir  = tmp_path / "custom_functions"
    custom_dir.mkdir()

    wb = Workbook()
    wb.remove(wb.active)
    items = wb.create_sheet("Items")
    items.append(["ItemId"])
    items.append(["A1"])

    vr = wb.create_sheet("ValidationRules")
    vr.append(["SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName",
               "ValidationType", "Severity", "Color", "Active",
               "ReferenceDataSheet", "ReferenceDataColumn", "CustomFunctionName"])
    vr.append([1, "R1", "Bad rule", "Items", "NoSuchColumn", "EMPTY", "ERROR", "RED", "Yes", "", "", ""])

    er = wb.create_sheet("EnrichmentRules")
    er.append(["SequenceNum", "RuleCode", "RuleName", "SheetName", "TargetColumn",
               "CustomFunctionName", "FunctionArguments", "Active", "Color"])

    wb.save(input_file)

    config_data = {
        "input_file": str(input_file),
        "output_file": str(output_file),
        "log_file": str(log_file),
        "custom_functions_directory": str(custom_dir),
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config_data))

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["erp_migration_data_validator.py", "--config", str(config_path)])

    exit_code = erp_migration_data_validator.main()

    assert exit_code == 1
    # Unlike a preflight failure, a Step 6 failure happens AFTER the
    # workbook was copied — the output file must exist and explain why.
    assert output_file.exists()

    result_wb = load_workbook(str(output_file))
    assert "Summary" in result_wb.sheetnames
    values = [
        c.value for row in result_wb["Summary"].iter_rows() for c in row if c.value is not None
    ]
    assert any("NoSuchColumn" in str(v) for v in values)
    assert any("TERMINATED" in str(v) for v in values)


def test_ai_connectivity_check_failure_terminates_before_copy(tmp_path, monkeypatch):
    """FS 3.4/3.8/8.1: a failed ai_connection_test_prompt check terminates
    processing immediately, same treatment as any other config error — the
    output file must not be created."""
    input_file  = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    log_file    = tmp_path / "run.log"
    custom_dir  = tmp_path / "custom_functions"
    custom_dir.mkdir()

    build_input_workbook(input_file)

    config_data = {
        "input_file": str(input_file),
        "output_file": str(output_file),
        "log_file": str(log_file),
        "custom_functions_directory": str(custom_dir),
        "ai_enabled": True,
        "ai_model": "fake-model",
        "ai_api_key": "fake-key",
        "ai_provider": "anthropic",
        "ai_provider_settings": {
            "anthropic": {"ai_endpoint": "https://example.test", "ai_api_version": "v1"}
        },
        "ai_connection_test_prompt": "Respond with OK.",
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config_data))

    def always_fails(config, prompt):
        raise RuntimeError("bad key")

    monkeypatch.setattr(ai_client, "_dispatch", always_fails)
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["erp_migration_data_validator.py", "--config", str(config_path)])

    exit_code = erp_migration_data_validator.main()

    assert exit_code == 1
    assert not output_file.exists()


def test_ai_connectivity_check_success_allows_run_to_continue(tmp_path, monkeypatch):
    input_file  = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    log_file    = tmp_path / "run.log"
    custom_dir  = tmp_path / "custom_functions"
    custom_dir.mkdir()

    build_input_workbook(input_file)

    config_data = {
        "input_file": str(input_file),
        "output_file": str(output_file),
        "log_file": str(log_file),
        "custom_functions_directory": str(custom_dir),
        "ai_enabled": True,
        "ai_model": "fake-model",
        "ai_api_key": "fake-key",
        "ai_provider": "anthropic",
        "ai_provider_settings": {
            "anthropic": {"ai_endpoint": "https://example.test", "ai_api_version": "v1"}
        },
        "ai_connection_test_prompt": "Respond with OK.",
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config_data))

    monkeypatch.setattr(ai_client, "_dispatch", lambda config, prompt: "OK")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["erp_migration_data_validator.py", "--config", str(config_path)])

    exit_code = erp_migration_data_validator.main()

    assert exit_code == 0
    assert output_file.exists()


def test_pipeline_fails_cleanly_on_preflight_error(tmp_path, monkeypatch):
    input_file  = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"
    log_file    = tmp_path / "run.log"
    custom_dir  = tmp_path / "custom_functions"
    custom_dir.mkdir()

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Items").append(["ItemId"])
    ef = wb.create_sheet("ExternalFiles")
    ef.append(["SequenceNum", "SheetName", "FilePath", "Format", "Active",
               "OriginalSheetName", "CSVDelimiter"])
    ef.append([1, "Stock", "/does/not/exist.csv", "csv", "Yes", "", ","])
    wb.save(input_file)

    config_data = {
        "input_file": str(input_file),
        "output_file": str(output_file),
        "log_file": str(log_file),
        "custom_functions_directory": str(custom_dir),
    }
    config_path = tmp_path / "config.yaml"
    config_path.write_text(yaml.safe_dump(config_data))

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["erp_migration_data_validator.py", "--config", str(config_path)])

    exit_code = erp_migration_data_validator.main()

    assert exit_code == 1
    # Output file must NOT be created — preflight errors must block copy (FS 8.2)
    assert not output_file.exists()
