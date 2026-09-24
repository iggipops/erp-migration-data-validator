"""
Tests for dimpolicyvalidation — Stock Balance Dimension Validation (FS Appendix B).

Organized in the three layers used across this suite:

  Layer 1 — metadata / config validation (check_metadata, REQUIRED_PARAMS)
  Layer 2 — runtime execution behavior   (the B.7 logic chain, step by step)
  Layer 3 — output correctness           (the two demo workbooks, B.8 messages)

followed by an explicit section drawing the soft-skip / hard-error /
partial-state distinctions, which are easy to conflate and carry very
different consequences for the run.

Layers 1 and 2 use small synthetic workbooks so each branch is isolated.
Layer 3 uses the two shipped demo workbooks and the two shipped config
files as ground truth, so the expected issue set is asserted against real
data and the configs themselves stay under test.
"""
import copy
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from config.config_loader import AppConfig
from importer.importer import Importer
from metadata.metadata_reader import MetadataReader
from functions.custom.dimpolicyvalidation import (
    REQUIRED_PARAMS, check_metadata, dimpolicyvalidation,
)
from registry.function_registry import FunctionRegistry
from validators.builtin import custom as custom_validator
from workbook.excel_utils import sheet_to_dataframe


REPO_ROOT  = Path(__file__).resolve().parents[4]
CUSTOM_DIR = REPO_ROOT / "src" / "v2" / "functions" / "custom"

DEMO_FILE = {
    "universal": REPO_ROOT / "test_data" / "v2" / "demo" / "dimpolicy_universal.xlsx",
    "ax":        REPO_ROOT / "test_data" / "v2" / "demo" / "dimpolicy_ax.xlsx",
}
DEMO_CONFIG = {
    "universal": REPO_ROOT / "src" / "v2" / "config_universal.example.yaml",
    "ax":        REPO_ROOT / "src" / "v2" / "config_ax.example.yaml",
}


# ---------------------------------------------------------------------------
# Synthetic fixtures (Layers 1 and 2)
# ---------------------------------------------------------------------------

STOCK_HEADER = [
    "Item ID", "Warehouse ID", "Location ID", "Lot Number", "Serial Number", "On-Hand",
]

# I1 Serial / I2 Lot / I3 NotUsed ; W1 Location Control Yes / W2 No
DEFAULT_STOCK = [
    ["I1", "W2", None,  None,    "SN-1", 1],   # Serial|No   -> passes
    ["I2", "W1", "L-1", "LOT-1", None,   7],   # Lot|Yes     -> passes
    ["I3", "W1", "L-2", None,    None,   7],   # NotUsed|Yes -> passes
]

DIM_POLICY_ROWS = [
    ["Key", "Lot Number", "Serial Number", "Location ID", "Quantity Restriction"],
    ["Serial|Yes",  "No",  "Yes", "Yes", 1],
    ["Serial|No",   "No",  "Yes", "No",  1],
    ["Lot|Yes",     "Yes", "No",  "Yes", None],
    ["Lot|No",      "Yes", "No",  "No",  None],
    ["NotUsed|Yes", "No",  "No",  "Yes", None],
    ["NotUsed|No",  "No",  "No",  "No",  None],
]


def build_workbook(make_workbook, sheet_factory, stock_rows=None, dim_policy_rows=None):
    wb = make_workbook()
    sheet_factory(wb, "ItemsMaster", [
        ["Item ID", "Tracking"],
        ["I1", "Serial"],
        ["I2", "Lot"],
        ["I3", "NotUsed"],
    ])
    sheet_factory(wb, "WarehousesMaster", [
        ["Warehouse ID", "Loc Control"],
        ["W1", "Yes"],
        ["W2", "No"],
    ])
    sheet_factory(wb, "DimPolicy", dim_policy_rows or DIM_POLICY_ROWS)
    sheet_factory(wb, "StockBalances",
                  [STOCK_HEADER] + (DEFAULT_STOCK if stock_rows is None else stock_rows))
    return wb


BASE_CFG = {
    "policy_key_lookups": [
        {"source_column": "Item ID", "reference_sheet": "ItemsMaster",
         "reference_key_column": "Item ID", "policy_key_columns": ["Tracking"]},
        {"source_column": "Warehouse ID", "reference_sheet": "WarehousesMaster",
         "reference_key_column": "Warehouse ID", "policy_key_columns": ["Loc Control"]},
    ],
    "dim_policy_key_delimiter": "|",
    "dim_policy_sheet": "DimPolicy",
    "dim_policy_key_column": "Key",
    "stock_to_dim_policy_columns_mapping": {
        "Lot Number": "Lot Number",
        "Serial Number": "Serial Number",
        "Location ID": "Location ID",
    },
    "stock_quantity_column": "On-Hand",
}


def cfg(**overrides):
    """A deep copy of BASE_CFG with top-level keys overridden."""
    c = copy.deepcopy(BASE_CFG)
    c.update(overrides)
    return c


def cfg_without(*keys):
    c = copy.deepcopy(BASE_CFG)
    for key in keys:
        c.pop(key, None)
    return c


def cfg_block(index, **overrides):
    """A deep copy of BASE_CFG with one lookup block's keys overridden."""
    c = copy.deepcopy(BASE_CFG)
    c["policy_key_lookups"][index].update(overrides)
    return c


def cfg_block_without(index, *keys):
    c = copy.deepcopy(BASE_CFG)
    for key in keys:
        c["policy_key_lookups"][index].pop(key, None)
    return c


class FakeConfig:
    def __init__(self, function_config):
        self.custom_functions = {"dimpolicyvalidation": function_config}


def errors_for(wb, function_config, sheet_name="StockBalances"):
    return check_metadata({"SheetName": sheet_name}, wb, FakeConfig(function_config))


def run(wb, function_config=None, sheet_name="StockBalances"):
    return dimpolicyvalidation({
        "worksheet_data":   sheet_to_dataframe(wb[sheet_name]),
        "workbook":         wb,
        "technical_config": BASE_CFG if function_config is None else function_config,
    })


def ids_of(errors):
    """The FS check ids ('B.5.1/A.1', ...) present in a list of error strings."""
    found = []
    for err in errors:
        text = err.rstrip()
        if text.endswith(")") and "(" in text:
            found.append(text[text.rfind("(") + 1:-1])
    return found


# ===========================================================================
# LAYER 1 — METADATA / CONFIG VALIDATION
# ===========================================================================

# --- B.4 / FS 7.2.1: declaration and registration --------------------------

def test_required_params_matches_spec_exactly():
    assert REQUIRED_PARAMS == [
        "policy_key_lookups",
        "dim_policy_sheet",
        "dim_policy_key_column",
        "stock_to_dim_policy_columns_mapping",
        "stock_quantity_column",
    ]


def test_registry_exposes_required_params_and_metadata_hook():
    reg = FunctionRegistry()
    reg.register_custom_directory(str(CUSTOM_DIR))

    assert reg.has("dimpolicyvalidation")
    assert reg.get_required_params("dimpolicyvalidation") == REQUIRED_PARAMS

    hook = reg.get_metadata_check_fn("dimpolicyvalidation")
    assert hook is not None and hook.__name__ == "check_metadata"


def test_module_registers_only_the_function_it_defines():
    """Imported helpers must never be registered as callable custom functions."""
    reg = FunctionRegistry()
    reg.register_custom_directory(str(CUSTOM_DIR))

    for helper in ("excel_utils", "sheet_exists", "get_headers",
                   "get_sheet", "sheet_to_dataframe", "find_sheet_name"):
        assert not reg.has(helper), f"{helper} must not be registered"


def test_metadata_validator_reports_missing_required_param(make_workbook, sheet_factory,
                                                           fake_config, vtype_registry):
    from metadata.metadata_validator import MetadataValidator

    wb = build_workbook(make_workbook, sheet_factory)
    sheet_factory(wb, "ValidationRules", [
        ["SequenceNum", "RuleCode", "RuleName", "SheetName", "ColumnName",
         "ValidationType", "Severity", "Color", "Active",
         "ReferenceDataSheet", "ReferenceDataColumn", "CustomFunctionName"],
        [1, "R1", "Dim policy", "StockBalances", "Item ID",
         "CUSTOM", "ERROR", "RED", "Yes", "", "", "dimpolicyvalidation"],
    ])
    sheet_factory(wb, "EnrichmentRules", [
        ["SequenceNum", "RuleCode", "RuleName", "SheetName", "TargetColumn",
         "CustomFunctionName", "FunctionArguments", "Active", "Color"],
    ])

    registry = FunctionRegistry()
    registry.register_custom_directory(str(CUSTOM_DIR))
    fake_config.custom_functions = {
        "dimpolicyvalidation": cfg_without("stock_quantity_column")
    }

    errors = MetadataValidator(wb, fake_config, registry, vtype_registry).validate()
    assert any("missing required parameter(s)" in e and "stock_quantity_column" in e
               for e in errors)


# --- B.5.1: schema validation, A.1-A.7 -------------------------------------

def test_a1_policy_key_lookups_absent(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    assert "B.5.1/A.1" in ids_of(errors_for(wb, cfg_without("policy_key_lookups")))


def test_a1_policy_key_lookups_empty_list(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    assert "B.5.1/A.1" in ids_of(errors_for(wb, cfg(policy_key_lookups=[])))


def test_a1_policy_key_lookups_not_a_list(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    bad = cfg(policy_key_lookups={"source_column": "Item ID"})
    assert "B.5.1/A.1" in ids_of(errors_for(wb, bad))


@pytest.mark.parametrize("missing_key", [
    "source_column", "reference_sheet", "reference_key_column", "policy_key_columns",
])
def test_a2_block_missing_a_required_key(make_workbook, sheet_factory, missing_key):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_block_without(0, missing_key))
    assert "B.5.1/A.2" in ids_of(errors)
    assert any(missing_key in e for e in errors)


def test_a2_block_is_not_a_mapping(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    bad = cfg(policy_key_lookups=["Item ID"])
    assert "B.5.1/A.2" in ids_of(errors_for(wb, bad))


def test_a3_policy_key_columns_empty_list(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    assert "B.5.1/A.3" in ids_of(errors_for(wb, cfg_block(0, policy_key_columns=[])))


def test_a3_policy_key_columns_as_bare_string(make_workbook, sheet_factory):
    """A single string is not a list — A.3 requires an ordered list."""
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_block(0, policy_key_columns="Tracking"))
    assert "B.5.1/A.3" in ids_of(errors)


@pytest.mark.parametrize("param", ["dim_policy_sheet", "dim_policy_key_column"])
def test_a4_required_scalars_absent(make_workbook, sheet_factory, param):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_without(param))
    assert "B.5.1/A.4" in ids_of(errors)
    assert any(param in e for e in errors)


@pytest.mark.parametrize("param", ["dim_policy_sheet", "dim_policy_key_column"])
def test_a4_required_scalars_blank(make_workbook, sheet_factory, param):
    wb = build_workbook(make_workbook, sheet_factory)
    assert "B.5.1/A.4" in ids_of(errors_for(wb, cfg(**{param: "   "})))


def test_a5_mapping_absent(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_without("stock_to_dim_policy_columns_mapping"))
    assert "B.5.1/A.5" in ids_of(errors)


def test_a5_mapping_empty(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg(stock_to_dim_policy_columns_mapping={}))
    assert "B.5.1/A.5" in ids_of(errors)


def test_a6_stock_quantity_column_absent(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_without("stock_quantity_column"))
    assert "B.5.1/A.6" in ids_of(errors)


def test_a7_delimiter_is_optional_and_defaults_to_empty_string(make_workbook, sheet_factory):
    """A.7: an absent delimiter is not an error — it defaults to ''."""
    wb = build_workbook(make_workbook, sheet_factory)
    assert errors_for(wb, cfg_without("dim_policy_key_delimiter")) == []


def test_a7_delimiter_must_be_a_scalar(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg(dim_policy_key_delimiter=["|", "-"]))
    assert "B.5.1/A.7" in ids_of(errors)


def test_schema_errors_are_collected_not_stopped_at_the_first(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    broken = cfg_without("policy_key_lookups", "dim_policy_sheet",
                         "stock_to_dim_policy_columns_mapping", "stock_quantity_column")
    found = set(ids_of(errors_for(wb, broken)))
    assert {"B.5.1/A.1", "B.5.1/A.4", "B.5.1/A.5", "B.5.1/A.6"} <= found


# --- B.5.2: cross-reference validation, B.1-B.8 ----------------------------

def test_b1_source_column_missing_from_stock_sheet(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_block(0, source_column="No Such Column"))
    assert "B.5.2/B.1" in ids_of(errors)


def test_b2_reference_sheet_missing(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_block(0, reference_sheet="NoSuchSheet"))
    assert "B.5.2/B.2" in ids_of(errors)


def test_b3_reference_key_column_missing(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_block(0, reference_key_column="No Such Column"))
    assert "B.5.2/B.3" in ids_of(errors)


def test_b4_policy_key_column_missing_from_reference_sheet(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_block(0, policy_key_columns=["Tracking", "Nope"]))
    assert "B.5.2/B.4" in ids_of(errors)


def test_b5_dim_policy_sheet_missing(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg(dim_policy_sheet="NoSuchSheet"))
    assert "B.5.2/B.5" in ids_of(errors)


def test_b6_dim_policy_key_column_missing(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg(dim_policy_key_column="No Such Column"))
    assert "B.5.2/B.6" in ids_of(errors)


def test_b7_mapping_key_missing_from_stock_sheet(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg(stock_to_dim_policy_columns_mapping={"Nope": "Lot Number"}))
    assert "B.5.2/B.7" in ids_of(errors)
    assert any("Nope" in e for e in errors)


def test_b7_mapping_value_missing_from_dim_policy_sheet(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg(stock_to_dim_policy_columns_mapping={"Lot Number": "Nope"}))
    assert "B.5.2/B.7" in ids_of(errors)


def test_b8_stock_quantity_column_missing(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg(stock_quantity_column="No Such Column"))
    assert "B.5.2/B.8" in ids_of(errors)


def test_cross_reference_errors_identify_the_offending_block(make_workbook, sheet_factory):
    """An error in the second block must say block 2, not block 1."""
    wb = build_workbook(make_workbook, sheet_factory)
    errors = errors_for(wb, cfg_block(1, reference_sheet="NoSuchSheet"))
    assert any("block 2" in e for e in errors)
    assert not any("block 1" in e for e in errors)


def test_cross_reference_errors_from_several_blocks_collected_together(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    broken = copy.deepcopy(BASE_CFG)
    broken["policy_key_lookups"][0]["source_column"]   = "Nope1"
    broken["policy_key_lookups"][1]["reference_sheet"] = "NoSuchSheet"
    errors = errors_for(wb, broken)
    assert any("block 1" in e for e in errors)
    assert any("block 2" in e for e in errors)


# --- B.5.3: data integrity validation, C.1-C.3 -----------------------------

def test_c1_duplicate_policy_key(make_workbook, sheet_factory):
    rows = DIM_POLICY_ROWS + [["Serial|No", "No", "Yes", "No", 1]]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    errors = errors_for(wb, BASE_CFG)
    assert "B.5.3/C.1" in ids_of(errors)
    assert any("Serial|No" in e for e in errors)


def test_c1_duplicate_detection_is_case_insensitive(make_workbook, sheet_factory):
    rows = DIM_POLICY_ROWS + [["serial|no", "No", "Yes", "No", 1]]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    assert "B.5.3/C.1" in ids_of(errors_for(wb, BASE_CFG))


def test_c1_empty_key_cells_are_not_duplicates_of_each_other(make_workbook, sheet_factory):
    rows = DIM_POLICY_ROWS + [[None, "No", "No", "No", None],
                              [None, "No", "No", "No", None]]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    assert "B.5.3/C.1" not in ids_of(errors_for(wb, BASE_CFG))


def test_c2_mapped_column_with_a_non_yes_no_value(make_workbook, sheet_factory):
    rows = DIM_POLICY_ROWS + [["Other|Yes", "Maybe", "No", "Yes", None]]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    errors = errors_for(wb, BASE_CFG)
    assert "B.5.3/C.2" in ids_of(errors)
    assert any("Maybe" in e and "Lot Number" in e for e in errors)


def test_c2_empty_cell_in_a_mapped_column_is_reported(make_workbook, sheet_factory):
    rows = DIM_POLICY_ROWS + [["Other|Yes", None, "No", "Yes", None]]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    errors = errors_for(wb, BASE_CFG)
    assert "B.5.3/C.2" in ids_of(errors)
    assert any("(empty)" in e for e in errors)


def test_c2_accepts_yes_no_case_and_whitespace_variants(make_workbook, sheet_factory):
    rows = DIM_POLICY_ROWS + [["Other|Yes", "yes", " NO ", "YES", None]]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    assert "B.5.3/C.2" not in ids_of(errors_for(wb, BASE_CFG))


def test_c3_non_numeric_quantity_restriction(make_workbook, sheet_factory):
    rows = DIM_POLICY_ROWS + [["Other|Yes", "No", "No", "Yes", "abc"]]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    errors = errors_for(wb, BASE_CFG)
    assert "B.5.3/C.3" in ids_of(errors)
    assert any("abc" in e for e in errors)


def test_c3_empty_quantity_restriction_cells_are_fine(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    assert "B.5.3/C.3" not in ids_of(errors_for(wb, BASE_CFG))


def test_c3_quantity_restriction_column_may_be_absent(make_workbook, sheet_factory):
    """FS B.6 makes the column optional — its absence is not an error."""
    rows = [row[:4] for row in DIM_POLICY_ROWS]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    assert errors_for(wb, BASE_CFG) == []


# --- Cross-layer -----------------------------------------------------------

def test_valid_synthetic_config_produces_no_errors(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    assert errors_for(wb, BASE_CFG) == []


@pytest.mark.parametrize("variant", ["universal", "ax"])
def test_shipped_config_is_valid_against_its_demo_workbook(variant):
    wb, _ = demo_results(variant)
    shipped = yaml.safe_load(DEMO_CONFIG[variant].read_text(encoding="utf-8"))
    function_config = shipped["custom_functions"]["dimpolicyvalidation"]
    assert errors_for(wb, function_config) == []


def test_config_names_resolve_case_insensitively(make_workbook, sheet_factory):
    """FS section 9 — sheet and column names are case-insensitive throughout."""
    wb = build_workbook(make_workbook, sheet_factory)
    mixed = {
        "policy_key_lookups": [
            {"source_column": "item id", "reference_sheet": "itemsmaster",
             "reference_key_column": "ITEM ID", "policy_key_columns": ["tracking"]},
            {"source_column": "WAREHOUSE ID", "reference_sheet": "WarehousesMASTER",
             "reference_key_column": "warehouse id", "policy_key_columns": ["LOC CONTROL"]},
        ],
        "dim_policy_key_delimiter": "|",
        "dim_policy_sheet": "dimpolicy",
        "dim_policy_key_column": "KEY",
        "stock_to_dim_policy_columns_mapping": {
            "lot number": "LOT NUMBER",
            "SERIAL NUMBER": "serial number",
            "location id": "Location ID",
        },
        "stock_quantity_column": "on-hand",
    }
    assert errors_for(wb, mixed) == []
    assert [passed for passed, _ in run(wb, mixed)] == [True, True, True]


def test_errors_from_all_three_layers_surface_together(make_workbook, sheet_factory):
    rows = DIM_POLICY_ROWS + [["Other|Yes", "Maybe", "No", "Yes", None]]
    wb = build_workbook(make_workbook, sheet_factory, dim_policy_rows=rows)
    broken = cfg_without("stock_quantity_column")           # A.6, schema
    broken["dim_policy_key_column"] = "No Such Column"      # B.6, cross-reference
    found = set(ids_of(errors_for(wb, broken)))             # C.2, data integrity
    assert "B.5.1/A.6" in found
    assert "B.5.2/B.6" in found
    assert "B.5.3/C.2" in found


# ===========================================================================
# LAYER 2 — RUNTIME EXECUTION BEHAVIOR (FS B.7)
# ===========================================================================

# --- Step 1a: per-block empty source value -> row skipped entirely ---------

def test_step1a_empty_source_in_first_block_skips_the_row(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[[None, "W2", "L-9", None, None, 99]])
    assert run(wb) == [(True, "")]


def test_step1a_empty_source_in_second_block_skips_the_row(make_workbook, sheet_factory):
    """Block 1 resolves fine; block 2's empty source still skips the whole row."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", None, "L-9", "LOT-X", None, 99]])
    assert run(wb) == [(True, "")]


@pytest.mark.parametrize("empty_value", [None, "", "   ", "nan", "NaN"])
def test_step1a_treats_blank_and_nan_as_empty(make_workbook, sheet_factory, empty_value):
    """FS 2.5.2 — 'nan' left by import/formula artifacts is empty, not text."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[[empty_value, "W2", "L-9", None, None, 99]])
    assert run(wb) == [(True, "")]


def test_step1a_skip_does_not_suppress_other_rows(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory, stock_rows=[
        [None, "W2", None,  None, "SN-1", 1],    # skipped
        ["I1", "W2", "L-9", None, "SN-2", 1],    # fails: Location ID must be absent
        ["I1", "W2", None,  None, "SN-3", 1],    # passes
    ])
    results = run(wb)
    assert results[0] == (True, "")
    assert results[1][0] is False
    assert results[2] == (True, "")


# --- Step 1b: unmatched source value -> terminal failure -------------------

def test_step1b_unmatched_value_fails_and_stops_further_checks(make_workbook, sheet_factory):
    """The row also violates the mapping, but only the 1b message is reported."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I9", "W2", "L-9", "LOT-X", "SN-X", 99]])
    passed, comment = run(wb)[0]
    assert passed is False
    assert comment == ("Item ID value I9 not found in ItemsMaster "
                       "— dimension policy cannot be determined")
    assert ";" not in comment


def test_step1b_names_the_failing_blocks_own_column_and_sheet(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W9", None, None, "SN-1", 1]])
    passed, comment = run(wb)[0]
    assert passed is False
    assert comment == ("Warehouse ID value W9 not found in WarehousesMaster "
                       "— dimension policy cannot be determined")


def test_step1b_in_first_block_short_circuits_before_second_block(make_workbook, sheet_factory):
    """Block 1 fails 1b; block 2's empty source must not turn this into a skip."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I9", None, None, None, None, 1]])
    passed, comment = run(wb)[0]
    assert passed is False
    assert "I9" in comment


# --- Steps 1c / 2: key composition ----------------------------------------

def test_key_components_join_in_block_order(make_workbook, sheet_factory):
    """I2 -> Lot, W1 -> Yes, so the key is 'Lot|Yes' and not 'Yes|Lot'."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I2", "W1", None, None, None, 1]])
    passed, comment = run(wb)[0]
    assert passed is False
    assert "policy Lot|Yes" in comment


def test_key_components_join_in_within_block_order(make_workbook, sheet_factory):
    """One block pulling two columns — the AX shape."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W1", None, None, None, 1]])
    sheet_factory(wb, "ItemsAX", [
        ["Item ID", "Tracking Group", "Storage Group"],
        ["I1", "Serial", "WH"],
    ])
    single = cfg(policy_key_lookups=[{
        "source_column": "Item ID", "reference_sheet": "ItemsAX",
        "reference_key_column": "Item ID",
        "policy_key_columns": ["Tracking Group", "Storage Group"],
    }])
    passed, comment = run(wb, single)[0]
    assert passed is False
    assert "Dimension policy key Serial|WH not found" in comment


def test_default_delimiter_concatenates_components_bare(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I2", "W1", None, None, None, 1]])
    passed, comment = run(wb, cfg_without("dim_policy_key_delimiter"))[0]
    assert passed is False
    assert "Dimension policy key LotYes not found" in comment


def test_multi_character_delimiter_is_honored(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I2", "W1", None, None, None, 1]])
    passed, comment = run(wb, cfg(dim_policy_key_delimiter="::"))[0]
    assert passed is False
    assert "Dimension policy key Lot::Yes not found" in comment


def test_composed_key_lookup_is_case_and_whitespace_insensitive(make_workbook, sheet_factory):
    rows = [
        ["Key", "Lot Number", "Serial Number", "Location ID", "Quantity Restriction"],
        ["  SERIAL|NO  ", "No", "Yes", "No", 1],
    ]
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", None, None, "SN-1", 1]],
                        dim_policy_rows=rows)
    assert run(wb) == [(True, "")]


# --- Step 3: composed key not found -> terminal failure --------------------

def test_step3_unknown_policy_key_fails_and_stops_further_checks(make_workbook, sheet_factory):
    rows = [r for r in DIM_POLICY_ROWS if r[0] != "Serial|No"]
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", "L-9", "LOT-X", None, 99]],
                        dim_policy_rows=rows)
    passed, comment = run(wb)[0]
    assert passed is False
    assert comment == "Dimension policy key Serial|No not found in Dim Policy sheet"
    assert ";" not in comment


# --- Step 4: mapped dimension columns, accumulating ------------------------

def test_step4_yes_with_empty_cell_fails(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I2", "W1", "L-1", None, None, 7]])
    passed, comment = run(wb)[0]
    assert passed is False
    assert comment == "Dimension column Lot Number is required by policy Lot|Yes but is empty"


def test_step4_yes_with_populated_cell_passes(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I2", "W1", "L-1", "LOT-1", None, 7]])
    assert run(wb) == [(True, "")]


def test_step4_no_with_populated_cell_fails(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", "L-9", None, "SN-1", 1]])
    passed, comment = run(wb)[0]
    assert passed is False
    assert comment == ("Dimension column Location ID must be absent for policy "
                       "Serial|No but value L-9 is present")


def test_step4_no_with_empty_cell_passes(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", None, None, "SN-1", 1]])
    assert run(wb) == [(True, "")]


def test_step4_reports_every_mismatch_in_mapping_order(make_workbook, sheet_factory):
    """Lot missing, Serial present, Location missing on one Lot|Yes row."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I2", "W1", None, None, "SN-X", 7]])
    passed, comment = run(wb)[0]
    assert passed is False
    parts = comment.split("; ")
    assert len(parts) == 3
    assert parts[0].startswith("Dimension column Lot Number is required")
    assert parts[1].startswith("Dimension column Serial Number must be absent")
    assert parts[2].startswith("Dimension column Location ID is required")


def test_step4_unparseable_policy_value_fails_the_row_and_accumulates():
    """FS B.5.3/C.2 makes a non-Yes/No policy cell a config-validation hard
    stop, so dimpolicyvalidation() itself rejects such a workbook before any
    row is read — this defensive branch is unreachable through the public
    entry point, which is why it is exercised on _validate_row directly.
    If config validation is ever bypassed, a corrupt policy table must still
    never pass silently, and the message must accumulate with the others."""
    from functions.custom.dimpolicyvalidation import _validate_row

    blocks = [{
        "source_column":      "Item ID",
        "reference_sheet":    "ItemsMaster",
        "lookup":             {"i1": {"Tracking": "Serial"}},
        "policy_key_columns": ["Tracking"],
    }]
    policy_lookup = {"serial": {"Lot Number": "Maybe", "Location ID": "No"}}
    mapping       = {"Lot Number": "Lot Number", "Location ID": "Location ID"}
    stock_row     = {"Item ID": "I1", "Lot Number": None,
                     "Location ID": "L-9", "On-Hand": 1}

    passed, comment = _validate_row(
        stock_row=stock_row, blocks=blocks, delimiter="",
        policy_lookup=policy_lookup, mapping=mapping,
        stock_qty_col="On-Hand", qty_restriction_col=None,
    )

    assert passed is False
    parts = comment.split("; ")
    assert len(parts) == 2
    assert parts[0] == ("Dimension column Lot Number has unparseable policy "
                        "value Maybe for policy Serial — must be Yes or No")
    assert parts[1].startswith("Dimension column Location ID must be absent")


def test_config_validation_rejects_a_non_yes_no_policy_table_up_front(
        make_workbook, sheet_factory):
    """The reachable half of the same concern: C.2 stops the run before any
    row is processed, so the defensive branch above should never fire."""
    rows = [
        ["Key", "Lot Number", "Serial Number", "Location ID", "Quantity Restriction"],
        ["Serial|No", "Maybe", "Yes", "No", None],
    ]
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", "L-9", None, "SN-1", 1]],
                        dim_policy_rows=rows)
    with pytest.raises(ValueError, match="B.5.3/C.2"):
        run(wb)


# --- Step 5: quantity restriction ------------------------------------------

def test_step5_quantity_mismatch_fails(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", None, None, "SN-1", 5]])
    passed, comment = run(wb)[0]
    assert passed is False
    assert comment == ("Quantity restriction for policy Serial|No requires 1 "
                       "but found 5 in column On-Hand")


def test_step5_matching_quantity_passes(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", None, None, "SN-1", 1]])
    assert run(wb) == [(True, "")]


def test_step5_skipped_when_restriction_is_empty(make_workbook, sheet_factory):
    """Lot|Yes carries no restriction, so any quantity is acceptable."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I2", "W1", "L-1", "LOT-1", None, 9999]])
    assert run(wb) == [(True, "")]


def test_step5_skipped_when_restriction_column_absent(make_workbook, sheet_factory):
    rows = [row[:4] for row in DIM_POLICY_ROWS]
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", None, None, "SN-1", 999]],
                        dim_policy_rows=rows)
    assert run(wb) == [(True, "")]


def test_step5_skipped_when_stock_quantity_is_empty(make_workbook, sheet_factory):
    """An empty quantity is left to a dedicated EMPTY/NULL rule, not reported here."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", None, None, "SN-1", None]])
    assert run(wb) == [(True, "")]


def test_step5_compares_numerically_across_types(make_workbook, sheet_factory):
    """Text '1' equals the number 1 — the restriction is a numeric comparison."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", None, None, "SN-1", "1"]])
    assert run(wb) == [(True, "")]


def test_step4_and_step5_failures_reported_together(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", "L-9", None, "SN-1", 5]])
    passed, comment = run(wb)[0]
    assert passed is False
    parts = comment.split("; ")
    assert len(parts) == 2
    assert parts[0].startswith("Dimension column Location ID must be absent")
    assert parts[1].startswith("Quantity restriction for policy Serial|No")


# --- Step 6 / function contract -------------------------------------------

def test_one_result_per_data_row_and_row_count_preserved(make_workbook, sheet_factory):
    stock = [["I1", "W2", None, None, f"SN-{i}", 1] for i in range(7)]
    wb = build_workbook(make_workbook, sheet_factory, stock_rows=stock)
    assert len(run(wb)) == 7


def test_fully_empty_row_still_yields_a_result(make_workbook, sheet_factory):
    """FS 2.5.1 — empty rows stay in scope and keep their position."""
    wb = build_workbook(make_workbook, sheet_factory, stock_rows=[
        ["I1", "W2", None, None, "SN-1", 1],
        [None, None, None, None, None, None],
        ["I1", "W2", None, None, "SN-2", 1],
    ])
    results = run(wb)
    assert len(results) == 3
    assert results[1] == (True, "")


def test_every_result_is_a_bool_and_comment_pair(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory, stock_rows=[
        ["I1", "W2", None,  None, "SN-1", 1],
        ["I1", "W2", "L-9", None, "SN-2", 1],
    ])
    for passed, comment in run(wb):
        assert isinstance(passed, bool)
        assert isinstance(comment, str)


def test_runtime_config_error_reports_every_problem_together(make_workbook, sheet_factory):
    wb = build_workbook(make_workbook, sheet_factory)
    broken = cfg_without("policy_key_lookups", "stock_to_dim_policy_columns_mapping")
    with pytest.raises(ValueError) as exc:
        run(wb, broken)
    message = str(exc.value)
    assert "B.5.1/A.1" in message
    assert "B.5.1/A.5" in message


@pytest.mark.parametrize("missing", ["worksheet_data", "workbook"])
def test_missing_context_element_raises(make_workbook, sheet_factory, missing):
    wb = build_workbook(make_workbook, sheet_factory)
    context = {
        "worksheet_data":   sheet_to_dataframe(wb["StockBalances"]),
        "workbook":         wb,
        "technical_config": BASE_CFG,
    }
    context[missing] = None
    with pytest.raises(ValueError):
        dimpolicyvalidation(context)


# ===========================================================================
# LAYER 3 — OUTPUT CORRECTNESS (the two demo workbooks, FS B.8)
# ===========================================================================

EM_DASH = "—"

EXPECTED = {
    "universal": {
        6:  "Quantity restriction for policy Serial|No requires 1 but found 12 in column On-Hand",
        13: "Dimension column Location ID must be absent for policy Serial|No "
            "but value E-01-01-1 is present",
        20: f"Item ID value Item_99 not found in ItemsMaster {EM_DASH} "
            f"dimension policy cannot be determined",
        26: "Quantity restriction for policy Serial|No requires 1 but found 2 in column On-Hand",
        34: f"Warehouse ID value WH4 not found in WarehousesMaster {EM_DASH} "
            f"dimension policy cannot be determined",
        35: "Dimension column Lot Number is required by policy Lot|Yes but is empty",
        63: "Dimension column Lot Number must be absent for policy NotUsed|Yes "
            "but value ILLEGAL-LOT-999 is present",
        86: "Dimension column Serial Number must be absent for policy Lot|Yes "
            "but value SN-ILLEGAL-88 is present",
        93: "Dimension column Location ID must be absent for policy Lot|No "
            "but value H-02-03-2 is present",
    },
    "ax": {
        6:  "Quantity restriction for policy Serial|WH requires 1 but found 12 in column On-Hand",
        13: "Dimension column Location ID must be absent for policy Serial|WH "
            "but value E-01-01-1 is present",
        20: f"Item ID value Item_99 not found in ItemsMaster {EM_DASH} "
            f"dimension policy cannot be determined",
        26: "Quantity restriction for policy Serial|WH requires 1 but found 2 in column On-Hand",
        29: "Quantity restriction for policy Serial|WH requires 1 but found 2 in column On-Hand",
        35: "Dimension column Lot Number is required by policy Lot|WH_Loc but is empty",
        63: "Dimension column Lot Number must be absent for policy NotUsed|WH_Loc "
            "but value ILLEGAL-LOT-999 is present",
        86: "Dimension column Serial Number must be absent for policy Lot|WH_Loc "
            "but value SN-ILLEGAL-88 is present",
    },
}


def demo_results(variant):
    """Run the shipped config against the shipped workbook, as the tool does.

    Mirrors erp_migration_data_validator.py Step 5 (FS 8.5): before validation can see
    ItemsMaster/WarehousesMaster/etc., the workbook's own ExternalFiles/
    ImportSpec sheets must be imported via the real Importer — the demo
    workbook ships in its pre-import shape, same as any real input file.
    """
    wb = load_workbook(DEMO_FILE[variant], data_only=True)
    shipped = yaml.safe_load(DEMO_CONFIG[variant].read_text(encoding="utf-8"))
    function_config = shipped["custom_functions"]["dimpolicyvalidation"]

    reader = MetadataReader(wb)
    external_files = reader.read_external_files()
    import_spec    = reader.read_import_spec()
    # FilePath entries are relative to the repo root (per the config files'
    # own header comments: "Run from the repository root") — make them
    # absolute so this resolves regardless of pytest's cwd.
    for rule in external_files:
        rule["FilePath"] = str(REPO_ROOT / rule["FilePath"])
        # The shipped demo workbooks predate ExternalFiles.CSVEncoding (FS 5.1);
        # every demo csv is plain ASCII, so utf-8 is the correct declaration.
        if str(rule.get("Format") or "").strip().lower() == "csv":
            rule.setdefault("CSVEncoding", "utf-8")

    dummy_config = AppConfig(
        input_file="", output_file="", log_file="", custom_functions_directory="",
    )
    Importer(wb, dummy_config).import_external_files(external_files, import_spec)

    return wb, run(wb, function_config)


@pytest.mark.parametrize("variant", ["universal", "ax"])
def test_demo_failing_rows_are_exactly_as_expected(variant):
    _, results = demo_results(variant)
    failing = {i + 1 for i, (passed, _) in enumerate(results) if not passed}
    assert failing == set(EXPECTED[variant])


@pytest.mark.parametrize("variant", ["universal", "ax"])
def test_demo_row_count_is_unchanged(variant):
    wb, results = demo_results(variant)
    assert len(results) == wb["StockBalances"].max_row - 1 == 100


@pytest.mark.parametrize(
    "variant,row",
    [("universal", r) for r in EXPECTED["universal"]] +
    [("ax", r) for r in EXPECTED["ax"]],
)
def test_demo_messages_match_b8_exactly(variant, row):
    _, results = demo_results(variant)
    passed, comment = results[row - 1]
    assert passed is False
    assert comment == EXPECTED[variant][row]


def test_row_93_fails_under_universal_but_passes_under_ax():
    """The clearest demonstration of what multi-block key composition buys.

    Item_29 is a Lot item sitting in WH1 with a WH3 location. The two-block
    key resolves Lot|No (WH1's Location Control is No) and flags the location
    that should not be there. The item-only AX key resolves Lot|WH_Loc, which
    requires a location, so the warehouse anomaly is invisible to it.
    """
    _, universal = demo_results("universal")
    _, ax        = demo_results("ax")

    assert universal[92][0] is False
    assert "Location ID must be absent for policy Lot|No" in universal[92][1]
    assert ax[92] == (True, "")


@pytest.mark.parametrize("variant,row,expected_key", [
    ("universal", 6,  "Serial|No"),
    ("universal", 35, "Lot|Yes"),
    ("universal", 63, "NotUsed|Yes"),
    ("universal", 93, "Lot|No"),
    ("ax", 6,  "Serial|WH"),
    ("ax", 35, "Lot|WH_Loc"),
    ("ax", 63, "NotUsed|WH_Loc"),
])
def test_demo_rows_resolve_to_the_expected_policy_key(variant, row, expected_key):
    _, results = demo_results(variant)
    assert f"policy {expected_key} " in results[row - 1][1]


@pytest.mark.parametrize("variant", ["universal", "ax"])
def test_text_typed_quantity_without_a_restriction_passes(variant):
    """Row 78 stores On-Hand as the text '450.5'. Its NotUsed policy carries no
    Quantity Restriction, so step 5 exits before the value is ever read. Typing
    is a NUMERIC rule's concern, not this function's."""
    wb, results = demo_results(variant)
    assert wb["StockBalances"].cell(row=79, column=7).value == "450.5"
    assert results[77] == (True, "")


@pytest.mark.parametrize("variant", ["universal", "ax"])
def test_b8_messages_carry_no_placeholder_brackets_or_added_quotes(variant):
    _, results = demo_results(variant)
    for passed, comment in results:
        if passed:
            continue
        assert "[" not in comment and "]" not in comment
        assert "'" not in comment and '"' not in comment


def test_not_found_message_uses_an_em_dash():
    _, results = demo_results("universal")
    assert EM_DASH in results[19][1]
    assert " - " not in results[19][1]


def test_end_to_end_through_the_custom_validator(fake_config, issues_writer_factory):
    """Comments reach the Issues sheet and the anchor cell is highlighted."""
    wb, _ = demo_results("universal")
    shipped = yaml.safe_load(DEMO_CONFIG["universal"].read_text(encoding="utf-8"))
    fake_config.custom_functions = {
        "dimpolicyvalidation": shipped["custom_functions"]["dimpolicyvalidation"]
    }

    registry = FunctionRegistry()
    registry.register_custom_directory(str(CUSTOM_DIR))
    issues_writer = issues_writer_factory(wb, fake_config)

    sheet = wb["StockBalances"]
    rule  = {"RuleCode": "DIMPOLICYVALIDATION", "RuleName": "Dim policy",
             "Severity": "ERROR", "Color": "RED",
             "CustomFunctionName": "dimpolicyvalidation"}

    custom_validator.validate(sheet, 1, "Item ID", rule, {
        "issues_writer": issues_writer, "fn_registry": registry,
        "config": fake_config, "workbook": wb,
    })

    assert issues_writer.issue_counter == 9

    issues   = wb["Issues"]
    recorded = {row[3]: row[8] for row in issues.iter_rows(min_row=2, values_only=True)}
    # Issues RowNumber is the worksheet row: data row + 1 for the header
    assert recorded[7]  == EXPECTED["universal"][6]
    assert recorded[94] == EXPECTED["universal"][93]
    assert sheet.cell(row=7, column=1).fill.start_color.rgb is not None


# ===========================================================================
# SOFT SKIP vs HARD ERROR vs PARTIAL STATE
# ===========================================================================

def record_issues(wb, function_config, fake_config, issues_writer_factory):
    fake_config.custom_functions = {"dimpolicyvalidation": function_config}
    registry = FunctionRegistry()
    registry.register_custom_directory(str(CUSTOM_DIR))
    issues_writer = issues_writer_factory(wb, fake_config)

    rule = {"RuleCode": "R1", "RuleName": "Dim policy", "Severity": "ERROR",
            "Color": "RED", "CustomFunctionName": "dimpolicyvalidation"}
    custom_validator.validate(wb["StockBalances"], 1, "Item ID", rule, {
        "issues_writer": issues_writer, "fn_registry": registry,
        "config": fake_config, "workbook": wb,
    })
    return issues_writer


def test_soft_skip_passes_and_writes_no_issue(make_workbook, sheet_factory,
                                              fake_config, issues_writer_factory):
    """Step 1a is a pass, not a suppressed failure — no Issues row is written,
    even though the row violates the mapping in three separate ways."""
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[[None, "W2", "L-9", "LOT-X", "SN-X", 99]])
    writer = record_issues(wb, BASE_CFG, fake_config, issues_writer_factory)
    assert writer.issue_counter == 0


def test_hard_row_failure_writes_one_issue(make_workbook, sheet_factory,
                                           fake_config, issues_writer_factory):
    wb = build_workbook(make_workbook, sheet_factory,
                        stock_rows=[["I1", "W2", "L-9", None, "SN-1", 1]])
    writer = record_issues(wb, BASE_CFG, fake_config, issues_writer_factory)
    assert writer.issue_counter == 1
    assert writer.error_counter == 1


def test_hard_rule_error_raises_so_the_engine_marks_the_rule_failed(make_workbook, sheet_factory):
    """FS 8.8.2 — an invalid config fails the whole rule; no issues are recorded."""
    wb = build_workbook(make_workbook, sheet_factory)
    with pytest.raises(ValueError):
        run(wb, cfg_without("policy_key_lookups"))


def test_never_returns_none_so_the_rule_can_never_be_partial(make_workbook, sheet_factory):
    """Unlike batched AI functions (FS 6.8), this function has no unresolved
    state — every row gets a real verdict, so 'partial' is unreachable."""
    wb = build_workbook(make_workbook, sheet_factory, stock_rows=[
        [None, "W2", None,  None, "SN-1", 1],    # soft skip
        ["I9", "W2", None,  None, "SN-2", 1],    # step 1b failure
        ["I1", "W2", "L-9", None, "SN-3", 1],    # step 4 failure
        ["I1", "W2", None,  None, "SN-4", 1],    # pass
    ])
    results = run(wb)
    assert len(results) == 4
    assert all(r is not None for r in results)
    assert all(isinstance(r, tuple) and isinstance(r[0], bool) for r in results)


def test_terminal_failure_is_partial_row_evaluation_not_partial_rule_state(
        make_workbook, sheet_factory):
    """A 1b failure stops steps 4-5 for that row only; other rows are unaffected."""
    wb = build_workbook(make_workbook, sheet_factory, stock_rows=[
        ["I9", "W2", "L-9", "LOT-X", "SN-X", 99],   # 1b: one message only
        ["I1", "W2", "L-9", None,    "SN-1", 5],    # steps 4+5: two messages
    ])
    results = run(wb)
    assert results[0][1].count(";") == 0
    assert results[1][1].count(";") == 1
