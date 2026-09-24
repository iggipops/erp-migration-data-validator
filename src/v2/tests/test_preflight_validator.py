import pytest
from metadata.preflight_validator import PreflightValidator


def build_wb(make_workbook, sheet_factory, ef_rows=None, isp_rows=None,
             ef_headers=None, isp_headers=None):
    wb = make_workbook()
    if ef_rows is not None:
        headers = ef_headers or ["SequenceNum", "SheetName", "FilePath", "Format",
                                  "Active", "OriginalSheetName", "CSVDelimiter", "CSVEncoding"]
        sheet_factory(wb, "ExternalFiles", [headers] + ef_rows)
    if isp_rows is not None:
        headers = isp_headers or ["ExternalFileSheetName", "ColumnName", "TargetType",
                                   "Active", "DateFormats", "DecimalSeparator", "ThousandsSeparator"]
        sheet_factory(wb, "ImportSpec", [headers] + isp_rows)
    return wb


def test_no_sheets_present_is_valid(make_workbook, sheet_factory):
    wb = make_workbook()
    pv = PreflightValidator(wb, config=None)
    assert pv.validate() == []


def test_valid_external_files_row(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "items.csv"
    f.write_text("a,b\n1,2\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    assert pv.validate() == []


def test_inactive_row_not_checked(make_workbook, sheet_factory):
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", "/does/not/exist.csv", "csv", "No", "", ",", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    assert pv.validate() == []


def test_missing_active_value_is_error(make_workbook, sheet_factory):
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", "/tmp/x.csv", "csv", None, "", ",", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("Active is empty" in e for e in errors)


def test_duplicate_sequencenum(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"
    f.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"],
        [1, "Stock", str(f), "csv", "Yes", "", ",", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("SequenceNum" in e for e in errors)


def test_duplicate_sheetname(make_workbook, sheet_factory, tmp_path):
    f1 = tmp_path / "a.csv"; f1.write_text("a\n1\n")
    f2 = tmp_path / "b.csv"; f2.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f1), "csv", "Yes", "", ",", "utf-8"],
        [2, "Items", str(f2), "csv", "Yes", "", ",", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("duplicates the SheetName" in e for e in errors)


def test_sheetname_collides_with_existing_workbook_sheet(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Existing", str(f), "csv", "Yes", "", ",", "utf-8"],
    ])
    wb.create_sheet("Existing")
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("already exists in the input workbook" in e for e in errors)


def test_filepath_not_found(make_workbook, sheet_factory):
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", "/definitely/not/here.csv", "csv", "Yes", "", ",", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("FilePath not found" in e for e in errors)


def test_invalid_format(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.txt"; f.write_text("x")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "txt", "Yes", "", ",", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("Format must be xlsx or csv" in e for e in errors)


def test_xlsx_requires_originalsheetname(make_workbook, sheet_factory, tmp_path):
    from openpyxl import Workbook as WB
    f = tmp_path / "a.xlsx"
    extwb = WB(); extwb.active.append(["x"]); extwb.save(f)
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "xlsx", "Yes", "", "", ""],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("OriginalSheetName is required" in e for e in errors)


def test_xlsx_originalsheetname_must_exist_in_external_file(make_workbook, sheet_factory, tmp_path):
    from openpyxl import Workbook as WB
    f = tmp_path / "a.xlsx"
    extwb = WB(); extwb.active.title = "RealSheet"; extwb.active.append(["x"]); extwb.save(f)
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "xlsx", "Yes", "WrongSheet", "", ""],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("not found in external file" in e for e in errors)


def test_xlsx_valid_originalsheetname_passes(make_workbook, sheet_factory, tmp_path):
    from openpyxl import Workbook as WB
    f = tmp_path / "a.xlsx"
    extwb = WB(); extwb.active.title = "RealSheet"; extwb.active.append(["x"]); extwb.save(f)
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "xlsx", "Yes", "RealSheet", "", ""],
    ])
    pv = PreflightValidator(wb, config=None)
    assert pv.validate() == []


def test_csv_must_not_have_originalsheetname(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "csv", "Yes", "SomeSheet", ",", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("OriginalSheetName must be empty" in e for e in errors)


def test_csv_requires_delimiter(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "csv", "Yes", "", "", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("CSVDelimiter is required" in e for e in errors)


def test_csv_invalid_delimiter(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "csv", "Yes", "", "~", "utf-8"],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("not one of the supported delimiters" in e for e in errors)


def test_duplicate_filepath_and_originalsheetname(make_workbook, sheet_factory, tmp_path):
    from openpyxl import Workbook as WB
    f = tmp_path / "a.xlsx"
    extwb = WB(); extwb.active.title = "S1"; extwb.active.append(["x"])
    extwb.create_sheet("S2")
    extwb.save(f)
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "xlsx", "Yes", "S1", "", ""],
        [2, "Stock", str(f), "xlsx", "Yes", "S1", "", ""],
    ])
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("duplicate FilePath" in e for e in errors)


def test_mandatory_header_missing(make_workbook, sheet_factory):
    wb = build_wb(make_workbook, sheet_factory,
                   ef_rows=[[1, "Items", "/x", "csv", "Yes"]],
                   ef_headers=["SequenceNum", "SheetName", "FilePath", "Format", "Active"])
    # Remove a mandatory header manually by rewriting first row without "Active"
    ws = wb["ExternalFiles"]
    ws.cell(row=1, column=5, value="NotActive")
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any('mandatory column "Active" not found' in e for e in errors)


def test_headers_matched_case_insensitively(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"]],
        ef_headers=["sequencenum", "sheetname", "filepath", "format", "active",
                    "originalsheetname", "csvdelimiter", "csvencoding"],
    )
    pv = PreflightValidator(wb, config=None)
    assert pv.validate() == []


# --- ImportSpec ---

def test_importspec_referential_integrity(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"]],
        isp_rows=[["NoSuchSheet", "Qty", "numeric", "Yes", "", "", ""]],
    )
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("does not reference an Active=Yes row" in e for e in errors)


def test_importspec_valid_reference_passes(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"]],
        isp_rows=[["Items", "Qty", "numeric", "Yes", "", "", ""]],
    )
    pv = PreflightValidator(wb, config=None)
    assert pv.validate() == []


def test_importspec_must_not_target_xlsx_external_file(make_workbook, sheet_factory, tmp_path):
    from openpyxl import Workbook
    f = tmp_path / "a.xlsx"
    ext = Workbook()
    ext.active.title = "Data"
    ext.save(f)
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "xlsx", "Yes", "Data", "", ""]],
        isp_rows=[["Items", "Qty", "numeric", "Yes", "", "", ""]],
    )
    errors = PreflightValidator(wb, config=None).validate()
    assert len(errors) == 1
    assert "ImportSpec row 2" in errors[0] and "csv external files only" in errors[0]


def test_importspec_on_inactive_row_may_name_xlsx_sheet(make_workbook, sheet_factory, tmp_path):
    from openpyxl import Workbook
    f = tmp_path / "a.xlsx"
    ext = Workbook()
    ext.active.title = "Data"
    ext.save(f)
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "xlsx", "Yes", "Data", "", ""]],
        isp_rows=[["Items", "Qty", "numeric", "No", "", "", ""]],
    )
    assert PreflightValidator(wb, config=None).validate() == []


def test_importspec_date_requires_dateformats(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"]],
        isp_rows=[["Items", "ShipDate", "date", "Yes", "", "", ""]],
    )
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("DateFormats is required" in e for e in errors)


def test_importspec_invalid_targettype(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"]],
        isp_rows=[["Items", "Qty", "float", "Yes", "", "", ""]],
    )
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("not supported" in e for e in errors)


def test_importspec_decimal_thousands_must_differ(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"]],
        isp_rows=[["Items", "Qty", "numeric", "Yes", "", ".", "."]],
    )
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("must differ" in e for e in errors)


def test_importspec_separator_must_be_single_char(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"]],
        isp_rows=[["Items", "Qty", "numeric", "Yes", "", "..", ""]],
    )
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("single character" in e for e in errors)


def test_importspec_duplicate_sheet_column_combo(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ",", "utf-8"]],
        isp_rows=[
            ["Items", "Qty", "numeric", "Yes", "", "", ""],
            ["Items", "Qty", "text", "Yes", "", "", ""],
        ],
    )
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("duplicate ExternalFileSheetName" in e for e in errors)


def test_shared_error_list_regression(make_workbook, sheet_factory, tmp_path):
    """Regression: an ExternalFiles row-level error must not suppress
    ImportSpec's own header-presence check / row checks."""
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", "/does/not/exist.csv", "csv", "Yes", "", ",", "utf-8"]],
        isp_rows=[["NoSuchSheet", "Qty", "numeric", "Yes", "", "", ""]],
    )
    pv = PreflightValidator(wb, config=None)
    errors = pv.validate()
    assert any("FilePath not found" in e for e in errors)
    assert any("does not reference an Active=Yes row" in e for e in errors)


# --- CSVEncoding (FS 5.1, 8.3.1) ---

def test_csv_requires_encoding(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "csv", "Yes", "", ",", ""],
    ])
    errors = PreflightValidator(wb, config=None).validate()
    assert any("row 2: CSVEncoding is required" in e for e in errors)


def test_csv_invalid_encoding(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "csv", "Yes", "", ",", "latin-1"],
    ])
    errors = PreflightValidator(wb, config=None).validate()
    assert any('CSVEncoding "latin-1" is not one of the supported encodings' in e for e in errors)


@pytest.mark.parametrize("enc", ["utf-8", "utf-8-sig", "cp1251", "cp1252",
                                 "UTF-8", "Utf-8-Sig", "CP1251", "CP1252"])
def test_csv_supported_encodings_accepted_case_insensitively(make_workbook, sheet_factory, tmp_path, enc):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "csv", "Yes", "", ",", enc],
    ])
    assert PreflightValidator(wb, config=None).validate() == []


def test_xlsx_must_not_have_encoding(make_workbook, sheet_factory, tmp_path):
    from openpyxl import Workbook as WB
    f = tmp_path / "a.xlsx"
    extwb = WB(); extwb.active.title = "Data"; extwb.active.append(["x"]); extwb.save(f)
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", str(f), "xlsx", "Yes", "Data", "", "utf-8"],
    ])
    errors = PreflightValidator(wb, config=None).validate()
    assert any("row 2: CSVEncoding must be empty when Format = xlsx" in e for e in errors)


def test_inactive_row_encoding_not_checked(make_workbook, sheet_factory):
    wb = build_wb(make_workbook, sheet_factory, ef_rows=[
        [1, "Items", "/does/not/exist.csv", "csv", "No", "", ",", "bogus"],
    ])
    assert PreflightValidator(wb, config=None).validate() == []


def test_missing_csvencoding_header_is_error(make_workbook, sheet_factory, tmp_path):
    f = tmp_path / "a.csv"; f.write_text("a\n1\n")
    wb = build_wb(
        make_workbook, sheet_factory,
        ef_rows=[[1, "Items", str(f), "csv", "Yes", "", ","]],
        ef_headers=["SequenceNum", "SheetName", "FilePath", "Format", "Active",
                    "OriginalSheetName", "CSVDelimiter"],
    )
    errors = PreflightValidator(wb, config=None).validate()
    assert any('column "CSVEncoding" not found' in e for e in errors)
